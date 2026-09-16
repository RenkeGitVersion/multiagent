from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph


DOCX = Path("/Users/renke/毕设/中期答辩/任科中期报告_0701_新框架版.docx")
OUT_IMG = Path("/Users/renke/Documents/Codex/2026-05-27/1-1-2-agent-1-agent/tmp/research_framework_0705.png")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/Supplemental/Songti.ttc",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size=size, index=1 if bold else 0)
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for raw in text.split("\n"):
        line = ""
        for ch in raw:
            test = line + ch
            if draw.textbbox((0, 0), test, font=fnt)[2] <= max_width:
                line = test
            else:
                if line:
                    lines.append(line)
                line = ch
        if line:
            lines.append(line)
    return lines


def rounded_box(draw, xy, fill, outline="#30343b", width=3, radius=26):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def text_center(draw, box, title, body="", fill="#111827"):
    x1, y1, x2, y2 = box
    title_font = font(34, True)
    body_font = font(24)
    title_lines = wrap_text(draw, title, title_font, x2 - x1 - 44)
    body_lines = wrap_text(draw, body, body_font, x2 - x1 - 54) if body else []
    total_h = len(title_lines) * 44 + (14 if body_lines else 0) + len(body_lines) * 33
    y = y1 + ((y2 - y1 - total_h) // 2)
    for line in title_lines:
        w = draw.textbbox((0, 0), line, font=title_font)[2]
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=title_font, fill=fill)
        y += 44
    y += 8
    for line in body_lines:
        w = draw.textbbox((0, 0), line, font=body_font)[2]
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=body_font, fill="#374151")
        y += 33


def arrow(draw, start, end, color="#3f4b5f"):
    draw.line([start, end], fill=color, width=5)
    sx, sy = start
    ex, ey = end
    if ey > sy:
        pts = [(ex, ey), (ex - 13, ey - 22), (ex + 13, ey - 22)]
    else:
        pts = [(ex, ey), (ex - 13, ey + 22), (ex + 13, ey + 22)]
    draw.polygon(pts, fill=color)


def draw_framework() -> None:
    OUT_IMG.parent.mkdir(parents=True, exist_ok=True)
    W, H = 1900, 2550
    img = Image.new("RGB", (W, H), "#f8fafc")
    draw = ImageDraw.Draw(img)
    title_f = font(54, True)
    sub_f = font(26)
    draw.text((W / 2, 70), "研究框架图", anchor="mm", font=title_f, fill="#111827")
    draw.text((W / 2, 128), "面向需求多样性的 AI 语音智能体个性化匹配与交互设计研究", anchor="mm", font=sub_f, fill="#4b5563")

    # Left rail: problem and research questions.
    left_x, left_w = 110, 500
    rounded_box(draw, (left_x, 210, left_x + left_w, 620), "#fff7ed", "#c2410c")
    text_center(
        draw,
        (left_x, 210, left_x + left_w, 620),
        "第1章\n提出问题",
        "家庭多用户多场景需求多样\n单一智能体适配不足\n多智能体带来发现与理解问题",
        "#7c2d12",
    )
    rounded_box(draw, (left_x, 690, left_x + left_w, 1010), "#eef2ff", "#4f46e5")
    text_center(
        draw,
        (left_x, 690, left_x + left_w, 1010),
        "研究问题",
        "RQ1 如何动态匹配合适智能体？\nRQ2 如何让多智能体交互可理解？",
        "#312e81",
    )
    arrow(draw, (left_x + left_w / 2, 620), (left_x + left_w / 2, 690))

    # Main vertical research flow.
    main_x, box_w = 720, 1060
    items = [
        ("第2章 相关研究与竞品分析", "文献研究：需求多样性、单智能体局限、多智能体优势\n竞品调研：传统单智能体音箱与新型多智能体音箱"),
        ("第3章 体验设计框架构建", "用户研究 → 用户体验重构 → 个性化匹配设计框架\n从唤醒 Agent 转向需求识别、Agent 匹配、多 Agent 交互、长期学习"),
        ("第4章 个性化匹配方法构建", "即时情境响应 → 用户偏好 → 长期记忆\n三层递进构成完整匹配链路"),
        ("第5章 多智能体交互设计研究", "外观与交互形象研究\n屏幕显示 × 信息共享实验\n形成可理解交互原则"),
        ("第6章 原型设计与开发", "系统设计与开发\n音箱设计与开发\n手机端控制台设计与开发"),
        ("第7章 实验验证", "原型用户体验验证\n方法有效性验证：个性化匹配 vs 唤醒词/手动切换"),
        ("第8章 研究总结", "理论贡献、方法贡献、设计贡献、实践贡献"),
    ]
    colors = [
        ("#ecfeff", "#0891b2", "#164e63"),
        ("#f0fdf4", "#16a34a", "#14532d"),
        ("#eff6ff", "#2563eb", "#1e3a8a"),
        ("#faf5ff", "#9333ea", "#581c87"),
        ("#fff7ed", "#ea580c", "#7c2d12"),
        ("#fef2f2", "#dc2626", "#7f1d1d"),
        ("#f8fafc", "#475569", "#111827"),
    ]
    y = 210
    h = 235
    gap = 72
    centers = []
    for idx, (title, body) in enumerate(items):
        fill, outline, tcolor = colors[idx]
        box = (main_x, y, main_x + box_w, y + h)
        rounded_box(draw, box, fill, outline, width=4, radius=28)
        text_center(draw, box, title, body, tcolor)
        centers.append((main_x + box_w / 2, y + h))
        if idx < len(items) - 1:
            arrow(draw, (main_x + box_w / 2, y + h), (main_x + box_w / 2, y + h + gap - 10))
        y += h + gap

    # Connector from RQ to main flow.
    draw.line([(left_x + left_w, 850), (main_x - 40, 850)], fill="#64748b", width=4)
    draw.polygon([(main_x - 40, 850), (main_x - 65, 837), (main_x - 65, 863)], fill="#64748b")

    # Bottom note.
    note = "逻辑主线：发现必要性 → 建立体验框架 → 构建匹配方法 → 设计可理解交互 → 开发原型 → 实验验证 → 总结贡献"
    draw.text((W / 2, H - 85), note, anchor="mm", font=font(28, True), fill="#334155")
    img.save(OUT_IMG, quality=95)


def insert_after(paragraph: Paragraph, text: str, style: str = "Report Body") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = style
    inserted.add_run(text)
    return inserted


def replace_exact(doc: Document, old: str, new: str) -> int:
    count = 0
    for paragraph in doc.paragraphs:
        if paragraph.text == old:
            paragraph.text = new
            count += 1
    return count


def delete_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def main() -> None:
    draw_framework()
    doc = Document(str(DOCX))

    # 2.2.1
    replace_exact(
        doc,
        "补充小米音箱等目前市面上的音箱的调研，强调他们的问题。",
        "传统单智能体音箱是目前家庭语音设备的主流形态，代表产品包括小米小爱音箱、天猫精灵、Apple HomePod、Amazon Echo 和 Google Nest Audio 等。这类产品通常以一个固定语音助手作为统一入口，围绕音乐播放、天气查询、闹钟提醒、智能家居控制、信息问答和简单任务管理展开。以小米小爱音箱为例，其优势在于与米家生态和大量家居设备联动，用户可以通过“小爱同学”完成灯光、空调、扫地机器人等设备控制；HomePod 更强调音质和苹果生态联动；Echo 和 Nest Audio 则分别依托 Alexa 与 Google Assistant 的服务能力。这些产品证明了语音音箱作为家庭入口的可行性，但它们的基本交互逻辑仍然是“一个助手服务所有需求”。",
    )
    p = None
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("传统单智能体音箱是目前家庭语音设备"):
            p = paragraph
            break
    if p is None:
        raise RuntimeError("2.2.1 replacement failed")
    p = insert_after(
        p,
        "从竞品体验看，传统单智能体音箱的问题不在于不能完成单项任务，而在于难以回应家庭生活中的需求差异。首先，单一助手通常以稳定的人设、音色和回答风格出现，面对儿童陪伴、健康咨询、学习辅导和家居控制等场景时，很难同时保持亲和、专业、可靠和高效。其次，传统音箱更擅长命令式任务，例如“打开客厅灯”“播放音乐”，但当用户提出“陪一下小孩”“我今天不太舒服”“帮我监督他看电视”这类带有关系、照护和情绪含义的需求时，系统往往只能按通用助手方式回答，难以自动选择更合适的角色和服务方式。",
    )
    insert_after(
        p,
        "再次，许多传统音箱虽然支持连续对话、儿童模式或家庭成员识别等功能，但这些能力通常仍然服务于同一个助手，并没有形成可被用户明确感知的多角色分工。用户也很难知道当前系统是否理解了自己的身份、场景和偏好。因此，传统单智能体音箱可以作为家庭语音入口的基础形态，却难以充分支持多用户、多场景下对人格、能力和交互风格的差异化需求。"
    )

    # 2.2.2
    old1 = "服务生态的局限性。不同语音智能体背后往往连接的是不同平台、能力和服务生态。例如 Alexa 更强于智能家居和电商，Google Assistant 更强于搜索与信息服务，Bixby 更贴近三星设备生态。对于家庭用户来说，真实需求并不会只落在一个生态里：同一台设备可能既要控制家电，又要查信息、播放内容、调用第三方服务或完成家庭任务。"
    new1 = "新型多智能体音箱试图突破传统单助手的服务边界，其核心思路是在同一台硬件设备中接入多个语音智能体或服务代理，让用户可以根据任务调用不同平台和能力。不同语音智能体背后往往连接的是不同服务生态，例如 Alexa 更强于智能家居和电商，Google Assistant 更强于搜索与信息服务，Bixby 更贴近三星设备生态。对于家庭用户来说，真实需求并不会只落在一个生态里：同一台设备可能既要控制家电，又要查信息、播放内容、调用第三方服务或完成家庭任务。因此，多智能体音箱在竞品层面的价值，是把单一生态入口扩展为多个服务能力共同存在的入口。"
    replace_exact(doc, old1, new1)

    old2 = "语音多智能体在服务生态上具有独特优势。亚马逊倡议，用户应当能够在同一设备上同时使用多个已注册的语音智能体，并通过各自唤醒词调用不同 agent。让一个硬件设备能够接入多个服务生态，使用户在不同任务中调用最合适的服务能力。亚马逊也发布了一份关于在单一设备上部署多个语音智能体的设计指南，该指南概述了三个主要领域的关键设计原则：隐私、可预测性和易用性[20]。"
    new2 = "亚马逊提出的 Voice Interoperability Initiative 可以视为这一方向的典型案例。它倡议用户应当能够在同一设备上同时使用多个已注册的语音智能体，并通过各自唤醒词调用不同 agent。亚马逊也发布了关于在单一设备上部署多个语音智能体的设计指南，强调隐私、可预测性和易用性，并要求设备向用户清晰表达聆听、思考、说话等注意力状态，以及让用户理解哪些录音或数据由哪个智能体处理[20]。这说明多智能体音箱不只是技术上接入多个助手，还必须处理身份识别、状态表达、隐私边界和信息共享等交互问题。"
    replace_exact(doc, old2, new2)

    target = None
    for paragraph in doc.paragraphs:
        if paragraph.text == new2:
            target = paragraph
            break
    if target is None:
        raise RuntimeError("2.2.2 replacement failed")
    p = insert_after(
        target,
        "但是，从竞品调研角度看，亚马逊方案仍然保留了较明显的使用门槛。第一，唤醒方式依赖多个唤醒词，用户需要记住不同智能体的名称或调用方式，也需要事先了解每个智能体擅长什么服务。对于老人、儿童或初次使用者来说，这会形成服务发现和记忆负担。第二，在对话过程中，用户不一定能清楚分辨当前是哪一个智能体正在聆听、思考或回答，尤其当多个 agent 共存时，如果视觉或声音状态提示不足，用户会难以判断当前交互对象和设备状态。",
    )
    insert_after(
        p,
        "第三，多智能体之间的会话连续性仍然是竞品中的难点。用户在一个智能体中说过的话，是否会被另一个智能体理解，哪些信息会被共享，什么时候应该保持当前对话，什么时候应该切换到其他智能体，这些问题如果只依赖唤醒词和后台规则，很难被用户感知和控制。因此，现有多智能体竞品提供了“多个服务生态接入同一设备”的重要方向，但还没有充分解决“如何自动匹配合适智能体”以及“如何让用户理解多智能体正在如何协作”这两个关键体验问题。"
    )

    # Replace research framework image paragraph.
    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "图1 研究框架图" and i > 0:
            img_para = doc.paragraphs[i - 1]
            img_para.clear()
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_para.add_run().add_picture(str(OUT_IMG), width=Inches(6.3))
            break
    else:
        raise RuntimeError("framework caption not found")

    doc.save(str(DOCX))
    print(DOCX)
    print(OUT_IMG)


if __name__ == "__main__":
    main()
