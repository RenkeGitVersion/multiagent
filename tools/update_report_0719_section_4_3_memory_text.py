from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOCX = Path("/Users/renke/毕设/中期答辩/任科中期报告_0719.docx")


def insert_after(paragraph: Paragraph, text: str, style: str = "Report Body") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = style
    inserted.add_run(text)
    return inserted


def find_exact(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"paragraph not found: {text}")


def find_exact_or_none(doc: Document, text: str) -> Paragraph | None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    return None


def replace_paragraph(paragraph: Paragraph, text: str, style: str = "Report Body") -> None:
    paragraph.clear()
    paragraph.style = style
    paragraph.add_run(text)


def main() -> None:
    doc = Document(str(DOCX))

    # Correct one explanatory sentence that was accidentally styled as a figure caption.
    old_memory_sentence = find_exact_or_none(
        doc,
        "用户的每条语录，会在本地进行加工，判断是否适合写入为长期记忆，或者说是否跟之前的记忆有冲突，需要更新。具体类型包括事实类记忆列表、偏好类记忆列表。",
    )
    if old_memory_sentence is not None:
        replace_paragraph(
            old_memory_sentence,
            "用户的每条语音输入不会直接原样写入长期记忆，而是先在本地进行整理和筛选。系统会判断这句话是否包含稳定事实、长期偏好或需要更新的信息；如果与既有记忆冲突，则进入更新或确认流程。由此形成事实类记忆和偏好类记忆两类内容。",
        )

    old_voice_sentence = find_exact_or_none(
        doc,
        "关于实现的技术细节，首先是如何基于用户声音输入，为每个用户创建一个专属的档案，记载跟这个用户有关的内容。",
    )
    if old_voice_sentence is not None:
        replace_paragraph(
            old_voice_sentence,
            "在技术实现上，长期记忆首先需要解决“这句话是谁说的”。系统通过声纹识别为家庭成员建立用户档案，把语音输入与具体用户绑定。这样后续记录偏好、任务和提醒时，系统就能区分是儿童、家长还是其他家庭成员的需求。",
        )

    anchors = [
        (
            "图 记忆的作用链路",
            "图中可以看到，记忆并不是独立存在的数据库，而是参与后续路由与回复生成的中间层。短期记忆提供当前对话状态，中期记忆保留近期任务和话题，长期记忆沉淀稳定画像与偏好，三者共同影响 Top3 候选智能体的召回和最终回复方式。",
        ),
        (
            "图 短期、中期、长期记忆的例子",
            "这三类记忆的差异主要在于生命周期。短期记忆服务本轮对话，例如当前正在和哪个智能体交谈；中期记忆服务一段时间内的任务，例如提醒、计划和最近话题；长期记忆服务持续使用，例如家庭角色、常用智能体和偏好风格。区分生命周期，可以避免把临时状态误写成长期偏好。",
        ),
        (
            "图 如何从声音记录用户身份",
            "声纹注册流程由前端录音、后端上传、模型提取声纹特征和本地档案更新组成。只有当音频质量和样本数量达到要求时，系统才会认为该用户身份可用；否则返回未确认状态，避免把噪声或他人的声音写入错误档案。",
        ),
        (
            "图 新用户识别与身份信息存储",
            "新用户档案不仅保存声纹特征，还会记录昵称、家庭角色、注册样本数、档案状态和最近更新时间。这些字段用于后续识别、调试和权限控制，也为手机端展示和管理家庭成员提供基础。",
        ),
        (
            "图 一次语音识别输出的身份结果",
            "一次识别结果会返回用户 ID、身份来源、置信度、音频质量和识别原因。系统只有在置信度和音频质量达标时才读取或写入该用户的长期记忆；当结果为 unknown 或 failed 时，只保留当前会话判断，不进入长期档案。",
        ),
        (
            "图 长期记忆中用户画像档案内容",
            "用户画像档案用于回答“这个人是谁”。其中 userId 和 familyRole 确定家庭身份，centroid 和 sampleCount 支持声纹识别，status 和 updatedAt 用于判断档案是否有效、是否需要重新采样。",
        ),
        (
            "图 单个用户的长期记忆内容",
            "单个用户的长期记忆进一步回答“这个人长期喜欢什么、需要什么”。facts 记录相对稳定的事实，preferences 记录可用于匹配的偏好，confidence 和 updatedAt 用于控制记忆是否可靠、是否需要更新。",
        ),
        (
            "图 异步任务的记忆内容",
            "异步任务记忆用于记录“未来要替用户完成什么”。其中 triggerAt 决定提醒时间，audience 决定提醒对象，message 决定提醒内容，reminderAgentId 决定由哪个智能体执行。这样系统不仅能记住用户偏好，也能在合适的时间用合适的角色完成后续提醒。",
        ),
    ]

    # Insert from bottom to top to keep paragraph references stable.
    for caption, body in reversed(anchors):
        paragraph = find_exact(doc, caption)
        next_para = paragraph._p.getnext()
        if next_para is not None:
            existing = Paragraph(next_para, paragraph._parent).text.strip()
            if existing == body:
                continue
        insert_after(paragraph, body)

    doc.save(str(DOCX))
    print(DOCX)


if __name__ == "__main__":
    main()
