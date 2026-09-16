from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


DOCX = Path("/Users/renke/毕设/中期答辩/任科中期报告_0701_新框架版.docx")
OUT_IMG = Path("/Users/renke/Documents/Codex/2026-05-27/1-1-2-agent-1-agent/tmp/research_framework_0705_refstyle.png")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/Supplemental/Songti.ttc",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size=size, index=1 if bold else 0)
    return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for raw in text.split("\n"):
        line = ""
        for ch in raw:
            test = line + ch
            if draw.textbbox((0, 0), test, font=fnt)[2] <= max_width:
                line = test
            else:
                if line:
                    lines.append(line)
                line = ch
        if line:
            lines.append(line)
    return lines


def centered_text(draw, xy, text, fnt, fill="#333333", line_gap=8):
    x1, y1, x2, y2 = xy
    lines = wrap(draw, text, fnt, x2 - x1 - 18)
    line_h = fnt.size + line_gap
    total_h = len(lines) * line_h - line_gap
    y = y1 + (y2 - y1 - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt)
        draw.text((x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2, y), line, font=fnt, fill=fill)
        y += line_h


def rect(draw, xy, fill="#ffffff", outline="#666666", width=3):
    draw.rectangle(xy, fill=fill, outline=outline, width=width)


def arrow_down(draw, x, y1, y2, color="#555555"):
    draw.line([(x, y1), (x, y2)], fill=color, width=5)
    draw.polygon([(x, y2 + 22), (x - 16, y2 - 4), (x + 16, y2 - 4)], fill=color)


def draw_framework() -> None:
    OUT_IMG.parent.mkdir(parents=True, exist_ok=True)
    W, H = 1700, 2350
    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)
    grey = "#d9d9d9"
    line = "#666666"
    title_f = font(38, True)
    h_f = font(34, True)
    cell_f = font(24)
    note_f = font(22)
    side_f = font(30, True)

    # Top title
    top = (350, 45, 1420, 105)
    rect(draw, top, grey, line, 3)
    centered_text(draw, top, "面向需求多样性的AI语音智能体个性化匹配与交互设计研究", title_f)
    arrow_down(draw, W // 2, 105, 160)

    sections = [
        {
            "side": "研究\n问题",
            "title": "一、研究背景与研究问题",
            "cells": ["需求多样性", "单一智能体不足", "多智能体方案", "交互需优化"],
            "note": "RQ1 动态匹配合适智能体    RQ2 语音交互可理解",
        },
        {
            "side": "体验\n框架",
            "title": "二、多智能体个性化体验设计框架",
            "cells": ["用户研究\n访谈+量表", "体验重构\n需求识别→匹配", "个性化匹配\n设计框架"],
            "note": "从唤醒Agent转向需求识别、Agent匹配、多Agent交互、长期学习",
        },
        {
            "side": "匹配\n方法",
            "title": "三、个性化匹配方法构建",
            "cells": ["Context Matching\n多模态识别+语义", "Preference Matching\n量表偏好数据", "Memory Matching\npersona/情境/临时上下文"],
            "note": "即时情境响应 → 用户偏好 → 长期记忆，逐层构成完整匹配链路",
        },
        {
            "side": "交互\n设计",
            "title": "四、多智能体交互设计研究",
            "cells": ["外观与形象研究\n显示能力+人格具象", "原型实验研究\n屏幕有无+信息共享", "设计原则\n可分辨/可感知/可控制"],
            "note": "从角色形象到信息边界，建立可理解的多智能体交互",
        },
        {
            "side": "原型\n设计",
            "title": "五、原型设计与开发",
            "cells": ["系统架构", "音箱外观", "音箱UI", "手机端APP"],
            "note": "形成可测试的带屏多智能体语音音箱与控制台",
        },
        {
            "side": "实验\n验证",
            "title": "六、实验验证",
            "cells": ["原型体验验证\n新原型 vs 小米", "方法有效性验证\n个性化匹配 vs 唤醒词", "访谈\n理解与信任"],
            "note": "验证体验价值、匹配有效性与交互可理解性",
        },
        {
            "side": "研究\n产出",
            "title": "七、研究产出总结",
            "cells": ["理论贡献\n体验设计框架", "方法贡献\n个性化匹配", "设计贡献\n交互原则", "实践贡献\n原型与验证"],
            "note": "形成面向需求多样性的多智能体语音交互研究产出",
        },
    ]

    x0, x1 = 270, 1475
    y = 160
    sec_h = 260
    gap = 42
    side_x = 82
    side_w = 70
    for idx, sec in enumerate(sections):
        # side label
        side_y = y + 36
        rect(draw, (side_x, side_y, side_x + side_w, side_y + 160), grey, line, 3)
        centered_text(draw, (side_x, side_y, side_x + side_w, side_y + 160), sec["side"], side_f)

        # outer box
        rect(draw, (x0, y, x1, y + sec_h), "#ffffff", line, 3)
        header = (615, y + 28, 1130, y + 82)
        rect(draw, header, grey, line, 3)
        centered_text(draw, header, sec["title"], h_f)

        # connector from header to cells
        mid = (header[0] + header[2]) // 2
        draw.line([(mid, header[3]), (mid, y + 112)], fill=line, width=3)

        cells = sec["cells"]
        n = len(cells)
        cell_gap = 30
        total_w = x1 - x0 - 140
        cell_w = (total_w - cell_gap * (n - 1)) / n
        cell_y1, cell_y2 = y + 112, y + 178
        first_x = x0 + 70
        draw.line([(first_x + cell_w / 2, y + 112), (first_x + (n - 1) * (cell_w + cell_gap) + cell_w / 2, y + 112)], fill=line, width=3)
        for i, c in enumerate(cells):
            cx1 = int(first_x + i * (cell_w + cell_gap))
            cx2 = int(cx1 + cell_w)
            draw.line([((cx1 + cx2) / 2, y + 112), ((cx1 + cx2) / 2, cell_y1)], fill=line, width=3)
            rect(draw, (cx1, cell_y1, cx2, cell_y2), "#ffffff", line, 3)
            centered_text(draw, (cx1, cell_y1, cx2, cell_y2), c, cell_f)

        note = (520, y + 188, 1225, y + 224)
        rect(draw, note, "#f2f2f2", line, 3)
        centered_text(draw, note, sec["note"], note_f)

        if idx < len(sections) - 1:
            arrow_down(draw, W // 2, y + sec_h, y + sec_h + gap - 24)
        y += sec_h + gap

    img.save(OUT_IMG, quality=95)


def replace_exact(doc: Document, old: str, new: str) -> int:
    count = 0
    for paragraph in doc.paragraphs:
        if paragraph.text == old:
            paragraph.text = new
            count += 1
    return count


def main() -> None:
    draw_framework()
    doc = Document(str(DOCX))

    replace_exact(
        doc,
        "亚马逊提出的 Voice Interoperability Initiative 可以视为这一方向的典型案例。它倡议用户应当能够在同一设备上同时使用多个已注册的语音智能体，并通过各自唤醒词调用不同 agent。亚马逊也发布了关于在单一设备上部署多个语音智能体的设计指南，强调隐私、可预测性和易用性，并要求设备向用户清晰表达聆听、思考、说话等注意力状态，以及让用户理解哪些录音或数据由哪个智能体处理[20]。这说明多智能体音箱不只是技术上接入多个助手，还必须处理身份识别、状态表达、隐私边界和信息共享等交互问题。",
        "亚马逊提出的 Voice Interoperability Initiative 可以视为这一方向的典型案例。它倡议用户应当能够在同一设备上同时使用多个已注册的语音智能体，并通过各自唤醒词调用不同 agent。需要注意的是，这一方案目前更接近“不同厂商语音助手的生态互通”：例如在同一硬件上并存两个不同厂家的语音智能体，让用户可以分别调用各自服务生态。其价值主要在于拓展服务生态和平台入口，而不是针对家庭中儿童、成人、老人、健康咨询、学习辅导、休闲陪伴等需求多样性建立一个可扩展的智能体角色池。亚马逊也发布了关于在单一设备上部署多个语音智能体的设计指南，强调隐私、可预测性和易用性，并要求设备向用户清晰表达聆听、思考、说话等注意力状态，以及让用户理解哪些录音或数据由哪个智能体处理[20]。这说明多智能体音箱不只是技术上接入多个助手，还必须处理身份识别、状态表达、隐私边界和信息共享等交互问题。",
    )

    replace_exact(
        doc,
        "但是，从竞品调研角度看，亚马逊方案仍然保留了较明显的使用门槛。第一，唤醒方式依赖多个唤醒词，用户需要记住不同智能体的名称或调用方式，也需要事先了解每个智能体擅长什么服务。对于老人、儿童或初次使用者来说，这会形成服务发现和记忆负担。第二，在对话过程中，用户不一定能清楚分辨当前是哪一个智能体正在聆听、思考或回答，尤其当多个 agent 共存时，如果视觉或声音状态提示不足，用户会难以判断当前交互对象和设备状态。",
        "但是，从竞品调研角度看，亚马逊方案仍然保留了较明显的使用门槛。第一，它目前主要解决“两个不同厂家的智能体如何在同一设备上共存”的生态问题，并没有进一步解决“面对不同用户和不同生活场景，系统如何主动选择合适智能体”的需求多样性问题。第二，唤醒方式仍然依赖多个唤醒词，用户需要记住不同智能体的名称或调用方式，也需要事先了解每个智能体擅长什么服务。对于老人、儿童或初次使用者来说，这会形成服务发现和记忆负担。第三，在对话过程中，用户不一定能清楚分辨当前是哪一个智能体正在聆听、思考或回答，尤其当多个 agent 共存时，如果视觉或声音状态提示不足，用户会难以判断当前交互对象和设备状态。",
    )

    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "图1 研究框架图" and i > 0:
            img_para = doc.paragraphs[i - 1]
            img_para.clear()
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_para.add_run().add_picture(str(OUT_IMG), width=Inches(5.8))
            break
    else:
        raise RuntimeError("framework image paragraph not found")

    doc.save(str(DOCX))
    print(DOCX)
    print(OUT_IMG)


if __name__ == "__main__":
    main()
