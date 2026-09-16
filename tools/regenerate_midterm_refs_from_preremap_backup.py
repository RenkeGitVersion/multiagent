from __future__ import annotations

import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SOURCE = Path("/Users/renke/毕设/中期答辩/任科中期报告_0621_最新_备份_引用重排前.docx")
OUTPUT = Path("/Users/renke/毕设/中期答辩/任科中期报告_0621_最新.docx")
PROPOSAL = Path("/Users/renke/毕设/开题报告最新_0107_任科_标文献.docx")

WANG_2025 = (
    "Wang, W., Lu, Z., Yuan, T., Fang, K., & Zhang, Y. (2025). "
    "In-vehicle multi-personality voice interaction: insights from Chinese automotive market "
    "on how dynamic personality influences driver affect and user experience. "
    "Ergonomics, 1–19. https://doi.org/10.1080/00140139.2025.2592984"
)


def iter_all_paragraphs(doc: Document):
    for para in doc.paragraphs:
        yield para
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield para


def insert_paragraph_before(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def copy_paragraph_format(src: Paragraph, dst: Paragraph) -> None:
    dst.paragraph_format.left_indent = src.paragraph_format.left_indent
    dst.paragraph_format.right_indent = src.paragraph_format.right_indent
    dst.paragraph_format.first_line_indent = src.paragraph_format.first_line_indent
    dst.paragraph_format.space_before = src.paragraph_format.space_before
    dst.paragraph_format.space_after = src.paragraph_format.space_after
    dst.paragraph_format.line_spacing = src.paragraph_format.line_spacing
    dst.paragraph_format.line_spacing_rule = src.paragraph_format.line_spacing_rule
    dst.alignment = src.alignment


def copy_run_format(src_run, dst_run) -> None:
    if src_run is None:
        return
    dst_run.bold = src_run.bold
    dst_run.italic = src_run.italic
    dst_run.underline = src_run.underline
    dst_run.font.name = src_run.font.name
    dst_run.font.size = src_run.font.size
    dst_run.font.color.rgb = src_run.font.color.rgb
    if src_run._r.rPr is not None:
        dst_run._r.get_or_add_rPr()
        dst_run._r.rPr.clear_content()
        for child in src_run._r.rPr:
            dst_run._r.rPr.append(deepcopy(child))


def extract_proposal_refs() -> dict[str, str]:
    doc = Document(str(PROPOSAL))
    start = None
    for i, para in enumerate(doc.paragraphs):
        if "参考文献" in para.text and "Reference" in para.text:
            start = i + 1
            break
    if start is None:
        raise RuntimeError("未找到开题报告中的参考文献标题")

    refs: dict[str, str] = {}
    index = 1
    for para in doc.paragraphs[start:]:
        text = para.text.strip()
        if not text:
            continue
        if "选题来源" in text or "拟解决的关键研究问题" in text:
            break
        refs[str(index)] = text
        index += 1

    refs["200"] = WANG_2025
    return refs


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    refs = extract_proposal_refs()

    shutil.copyfile(SOURCE, OUTPUT)
    doc = Document(str(OUTPUT))

    full_text = "\n".join(p.text for p in iter_all_paragraphs(doc))
    used_old_ids: list[str] = []
    for old_id in re.findall(r"\[(\d+)\]", full_text):
        if old_id not in used_old_ids:
            used_old_ids.append(old_id)

    missing = [old_id for old_id in used_old_ids if old_id not in refs]
    if missing:
        raise RuntimeError(f"开题报告参考文献中缺少这些编号: {missing}")

    id_map = {old_id: str(i) for i, old_id in enumerate(used_old_ids, start=1)}
    citation_re = re.compile(r"\[(\d+)\]")

    def repl(match: re.Match[str]) -> str:
        old_id = match.group(1)
        return f"[{id_map.get(old_id, old_id)}]"

    for para in iter_all_paragraphs(doc):
        for run in para.runs:
            if "[" in run.text and "]" in run.text:
                run.text = citation_re.sub(repl, run.text)

    # Remove any existing generated references block if present.
    paras = doc.paragraphs
    ref_start = None
    achievement_start = None
    for i, para in enumerate(paras):
        if para.text.strip() == "参考文献":
            ref_start = i
        if para.text.startswith("成果清单"):
            achievement_start = i
            break
    if ref_start is not None and achievement_start is not None and ref_start < achievement_start:
        for para in paras[ref_start:achievement_start]:
            para._element.getparent().remove(para._element)

    achievement_para = None
    for para in doc.paragraphs:
        if para.text.startswith("成果清单"):
            achievement_para = para
            break
    if achievement_para is None:
        raise RuntimeError("未找到“成果清单”段落，无法定位参考文献插入位置")

    body_template = next((p for p in doc.paragraphs if p.style.name == "Report Body"), None)
    body_style = "Report Body" if body_template else None
    heading_style = "Report Heading 1"

    insert_paragraph_before(achievement_para, "", body_style)
    insert_paragraph_before(achievement_para, "参考文献", heading_style)

    for old_id in used_old_ids:
        new_id = id_map[old_id]
        para = insert_paragraph_before(achievement_para, f"[{new_id}] {refs[old_id]}", body_style)
        if body_template:
            copy_paragraph_format(body_template, para)
            para.paragraph_format.first_line_indent = None
            para.paragraph_format.left_indent = None
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            if body_template.runs and para.runs:
                copy_run_format(body_template.runs[0], para.runs[0])

    doc.save(str(OUTPUT))

    print("引用编号映射:")
    for old_id in used_old_ids:
        print(f"[{old_id}] -> [{id_map[old_id]}] {refs[old_id][:110]}")


if __name__ == "__main__":
    main()
