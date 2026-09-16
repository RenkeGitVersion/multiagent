from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches


DOCX = Path("/Users/renke/毕设/中期答辩/任科中期报告_0629_新框架版.docx")
IMG = Path("/Users/renke/Documents/Codex/2026-05-27/1-1-2-agent-1-agent/tmp/research_framework_0630_memory.png")
TITLE = "面向需求多样性的AI语音智能体个性化匹配与交互设计研究"


def make_framework_image() -> None:
    from PIL import Image, ImageDraw, ImageFont

    IMG.parent.mkdir(parents=True, exist_ok=True)
    W, H = 1700, 2200
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    font_path = "/System/Library/Fonts/Hiragino Sans GB.ttc"
    title = ImageFont.truetype(font_path, 44)
    h1 = ImageFont.truetype(font_path, 31)
    small = ImageFont.truetype(font_path, 18)
    tiny = ImageFont.truetype(font_path, 16)

    def wrap(text: str, font, max_w: int) -> list[str]:
        lines, cur = [], ""
        for ch in text:
            if ch == "\n":
                lines.append(cur)
                cur = ""
                continue
            test = cur + ch
            if d.textbbox((0, 0), test, font=font)[2] <= max_w:
                cur = test
            else:
                if cur:
                    lines.append(cur)
                cur = ch
        if cur:
            lines.append(cur)
        return lines

    def centered_text(x, y, w, text, font, fill="#222222", line_gap=5):
        lines = wrap(text, font, int(w))
        total = len(lines) * font.size + (len(lines) - 1) * line_gap
        yy = y - total / 2
        for line in lines:
            box = d.textbbox((0, 0), line, font=font)
            d.text((x + (w - (box[2] - box[0])) / 2, yy), line, font=font, fill=fill)
            yy += font.size + line_gap

    def rect(x, y, w, h, fill="#ffffff", outline="#666666", width=2):
        d.rectangle((x, y, x + w, y + h), fill=fill, outline=outline, width=width)

    def line(x1, y1, x2, y2, width=3):
        d.line((x1, y1, x2, y2), fill="#555555", width=width)

    def arrow_down(x, y1, y2):
        line(x, y1, x, y2, 3)
        d.polygon([(x, y2 + 13), (x - 12, y2 - 6), (x + 12, y2 - 6)], fill="#555555")

    def box(text, x, y, w, h, fill="#ffffff", font=None):
        rect(x, y, w, h, fill, "#666666", 2)
        centered_text(x + 12, y + h / 2, w - 24, text, font or small)

    def vlabel(text, x, y, h):
        rect(x, y, 68, h, "#d9d9d9", "#666666", 2)
        yy = y + (h - len(text) * 33) / 2
        for ch in text:
            centered_text(x + 8, yy + 16, 52, ch, h1, fill="#222222", line_gap=0)
            yy += 33

    d.text((110, 54), f"研究框架：{TITLE}", font=title, fill="#111111")
    d.line((110, 124, 1590, 124), fill="#111111", width=6)
    main_w = 1180
    main_x = (W - main_w) // 2
    side_x = 88
    cx = W // 2
    y = 164
    box(TITLE, main_x + 55, y, main_w - 110, 66, "#d9d9d9", h1)
    arrow_down(cx, y + 66, y + 108)
    y += 118

    sections = [
        ("研究问题", "一、研究背景与研究问题", ["需求多样性", "单一智能体不足", "多智能体方案", "交互需优化"], "RQ1 动态匹配合适智能体    RQ2 语音交互可理解"),
        ("体验框架", "二、多智能体个性化体验设计框架", ["用户研究\n访谈+量表", "体验重构\n需求识别→匹配", "个性化匹配\n设计框架"], "即时情境响应 → 用户偏好 → 长期记忆"),
        ("匹配方法", "三、个性化匹配方法构建", ["Context Matching\n多模态识别+语义", "Preference Matching\n量表偏好数据", "Memory Matching\npersona/context/临时上下文"], "长期记忆同时承担记住信息与持续优化匹配"),
        ("交互设计", "四、多智能体交互设计研究", ["外观与形象研究\n显示能力+人格具象", "原型实验研究\n屏幕有无+信息共享", "设计原则\n可分辨/可感知/可控制"], "从角色形象到信息边界，建立可理解的多智能体交互"),
        ("原型设计", "五、原型设计", ["系统架构", "音箱外观", "音箱UI", "手机端APP"], "形成可测试的带屏多智能体语音音箱与控制台"),
        ("实验验证", "六、实验验证", ["原型体验验证\n新原型 vs 小米", "方法有效性验证\n个性化匹配 vs 唤醒词", "访谈\n理解与信任"], "验证体验价值、匹配有效性与交互可理解性"),
        ("研究产出", "七、研究产出总结", ["理论贡献", "方法贡献", "设计贡献", "实践贡献"], "沉淀多智能体语音交互体验框架、方法与原则"),
    ]
    for idx, (side, title_text, nodes, note) in enumerate(sections):
        sec_h = 232 if idx not in (2, 3) else 244
        sec_y = y
        vlabel(side, side_x, sec_y + 28, sec_h - 56)
        rect(main_x, sec_y, main_w, sec_h)
        box(title_text, main_x + 340, sec_y + 24, 500, 54, "#d9d9d9", h1)
        n = len(nodes)
        gap = 30
        bw = int((main_w - 140 - gap * (n - 1)) / n)
        start_x = main_x + 70
        line_y = sec_y + 108
        line(start_x + bw / 2, line_y, start_x + (n - 1) * (bw + gap) + bw / 2, line_y, 3)
        line(cx, sec_y + 78, cx, line_y, 3)
        for i, text in enumerate(nodes):
            x = start_x + i * (bw + gap)
            box(text, x, sec_y + 128, bw, 66, "#ffffff", small)
            line(x + bw / 2, line_y, x + bw / 2, sec_y + 128, 2)
        box(note, main_x + 250, sec_y + sec_h - 38, 680, 32, "#f2f2f2", tiny)
        if idx < len(sections) - 1:
            arrow_down(cx, sec_y + sec_h, sec_y + sec_h + 38)
        y += sec_h + 52
    img.save(IMG)


def set_para_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_paragraph_before(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = paragraph.__class__(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def add_before(anchor, typ: str, text: str = ""):
    style = {
        "h1": "Report Heading 1",
        "h2": "Report Heading 2",
        "h3": "Report Heading 3",
        "p": "Report Body",
    }[typ]
    return insert_paragraph_before(anchor, text, style)


def replace_range(doc: Document, start_idx: int, end_idx: int, blocks: list[tuple[str, str]]) -> None:
    anchor = doc.paragraphs[end_idx]
    for para in doc.paragraphs[start_idx:end_idx]:
        para._element.getparent().remove(para._element)
    for typ, text in blocks:
        add_before(anchor, typ, text)


def replace_picture(doc: Document) -> None:
    for i, para in enumerate(doc.paragraphs[:-1]):
        if para._p.xml.count("graphic") > 0 and doc.paragraphs[i + 1].text.strip() == "图1 研究框架图":
            para.clear()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.add_run().add_picture(str(IMG), width=Inches(6.65))
            return


def update_toc(doc: Document) -> None:
    replacements = {
        "3.2 用户偏好（Preference Matching）\n3.3 长期记忆（Memory Matching）\n3.4 持续个性化（Long-term Personalization）": "3.2 用户偏好（Preference Matching）\n3.3 长期记忆（Memory Matching）",
        "第4章 多智能体交互设计研究\n4.1 多智能体交互形象设计\n4.2 多模态状态反馈设计\n4.3 会话保持与智能体切换设计\n4.4 信息边界与用户控制": "第4章 多智能体交互设计研究\n4.1 研究定位：从匹配到可理解交互\n4.2 外观与交互形象研究\n4.3 屏幕显示与信息共享实验\n4.4 设计结论与原则",
    }
    for para in doc.paragraphs:
        if para.style.name == "Report TOC" and para.text.strip() in replacements:
            set_para_text(para, replacements[para.text.strip()])


def main() -> None:
    make_framework_image()
    doc = Document(str(DOCX))
    update_toc(doc)
    replace_picture(doc)

    idx = {p.text.strip(): i for i, p in enumerate(doc.paragraphs) if p.text.strip()}
    ch3 = idx["第3章 个性化匹配方法构建"]
    ch4 = idx["第4章 多智能体交互设计研究"]
    ch5 = idx["第5章 原型设计"]

    ch3_blocks = [
        ("h1", "第3章 个性化匹配方法构建"),
        ("p", "在体验框架的基础上，本研究将个性化匹配方法进一步拆分为三个层次：即时的情境响应、用户偏好和长期记忆。三个层次并不是彼此替代，而是从“当前这句话应该由谁回答”逐步扩展到“系统如何随着长期使用逐渐更懂这个家庭”。其中，原先单独表述的持续个性化不再作为独立概念，而是并入长期记忆：长期记忆既包括记住用户、话题和临时上下文，也包括根据这些记忆持续调整后续匹配策略。"),
        ("h2", "3.1 即时的情境响应（Context Matching）"),
        ("p", "即时的情境响应主要解决用户当下说出一句话时，系统如何快速判断合适的智能体。该层重点依据当前输入，而不是长期偏好。首先，系统通过语音或多模态方式识别用户的粗粒度特征，例如年龄段和性别，用于判断儿童、成人或中老年用户在同一 query 下可能需要的不同回应方式。其次，系统结合当前时间判断早、中、晚等使用情境，因为同样一句“陪我一下”在睡前、午后和早晨可能对应不同的陪伴方式。"),
        ("p", "在语义层面，系统使用大模型对用户 query 进行语义识别与匹配，判断其更接近休闲陪伴、健康咨询、智能家居、信息查询或学习辅导等场景。这里不只依赖关键词，而是关注用户真实需求。例如“我现在需要陪小孩”虽然没有直接说“讲故事”，但在儿童陪伴语境中更可能需要儿童友好的智能体，而不是健康顾问或通用助手。"),
        ("p", "即时情境响应还需要区分强意图和弱意图。强意图是用户明确说出智能体名称、别名或唤醒词，这类请求直接路由；弱意图是用户自然表达需求但没有指定智能体，需要系统结合多模态特征、当前时间和语义判断完成初步匹配。"),
        ("h2", "3.2 用户偏好（Preference Matching）"),
        ("p", "用户偏好匹配用于把前序用户研究和量表数据真正转化为系统行为。本研究在量表中将智能体特征归纳为亲密陪伴型、鼓励赞赏型、成熟关怀型、专业可靠型和中性工具型，并获得不同用户类型、不同场景下对这些特征的偏好分。系统在弱意图匹配时，会先根据用户类型和场景查找偏好分，再结合当前智能体池中每个智能体的特征标签进行排序。"),
        ("p", "具体而言，每个智能体可以被标注为一种或两种偏好类型。例如儿童陪伴智能体可能对应亲密陪伴型和鼓励赞赏型，健康顾问可能对应专业可靠型和成熟关怀型。系统根据用户类型与场景得到各类智能体特征的偏好分，再对智能体标签取平均或加权计算，得到候选智能体排序。该步骤通常召回 Top3 候选，再交给大模型结合 query、上下文和智能体描述做最终选择。"),
        ("p", "这一层的意义在于避免模型完全凭语义直觉选择智能体，而是引入用户研究得到的先验偏好。它使匹配逻辑能够解释为：在某类用户、某类场景中，用户更偏好什么样的智能体特征，因此系统优先推荐与这些特征更匹配的智能体。"),
        ("h2", "3.3 长期记忆（Memory Matching）"),
        ("p", "长期记忆匹配用于解决一次性情境判断无法覆盖的长期家庭使用问题。它包含两层含义：第一，系统需要记住用户画像、近期话题和临时上下文；第二，系统需要根据这些记忆持续调整后续匹配策略。因此，本研究不再将“持续个性化”单独拆出，而是将其视为长期记忆发挥作用后的结果。"),
        ("p", "本研究将长期记忆分为三类：persona、context 和临时上下文。persona 指相对稳定的用户画像，包括家庭成员身份、年龄段、常用智能体、偏好的语气、拒绝的风格、常见场景和隐私边界。这类信息更新频率较低，应当通过明确反馈、长期统计或用户在手机端设置后更新，避免系统因为一次偶然对话就改变用户画像。"),
        ("p", "context 指最近一段时间内的对话主题和任务状态，例如最近在聊英语学习、正在进行健康咨询、刚刚设置了看电视提醒等。这类信息更新频率中等，随着会话结束、任务完成或时间过去而衰减。临时上下文则指当前一轮对话中的短时信息，例如上一句话、当前智能体、是否处于会话保持期、用户是否刚刚打断。这类信息更新频率最高，通常只在当前会话窗口内有效。"),
        ("p", "长期记忆的价值不只是“存储信息”，而是改变后续匹配行为。例如，系统记住某个儿童更喜欢虎子讲故事，之后在儿童休闲陪伴场景中默认优先推荐虎子，并减少重复询问；系统记住某个家长希望作业提醒由温和的儿童向智能体完成，之后在类似任务中优先选择儿童友好型智能体。通过这种方式，个性化匹配从单轮对话走向长期家庭使用。"),
    ]
    replace_range(doc, ch3, ch4, ch3_blocks)

    idx = {p.text.strip(): i for i, p in enumerate(doc.paragraphs) if p.text.strip()}
    ch4 = idx["第4章 多智能体交互设计研究"]
    ch5 = idx["第5章 原型设计"]
    ch4_blocks = [
        ("h1", "第4章 多智能体交互设计研究"),
        ("p", "前文的个性化匹配方法解决的是“系统应该选择哪个智能体”。但在真实交互中，仅仅在后台完成匹配并不足够，用户还需要理解当前是谁在服务、系统为什么切换、哪些信息被不同智能体使用。因此，本章从交互设计角度展开，围绕多智能体语音音箱的“可理解性”进行研究。"),
        ("p", "本章包含两个互相关联的研究。第一个研究关注多智能体音箱的外观与交互形象，即音箱是否需要屏幕、智能体应以何种具象程度呈现、不同角色如何被用户感知。第二个研究关注原型交互实验，即在多智能体语音对话中，屏幕有无和智能体间是否共享信息会如何影响可用性、信任、认知负荷、隐私感知和智能体识别效率。两个研究共同服务于 RQ2：如何让多智能体语音交互可理解。"),
        ("h2", "4.1 研究定位：从匹配到可理解交互"),
        ("p", "多智能体系统如果只在后台完成匹配，用户仍然会觉得它像黑箱。与手机或电脑界面不同，智能音箱以语音为主，用户通常无法同时看到多个智能体并列工作。因此，当多个角色通过同一个设备说话时，用户容易产生三个疑问：当前是谁在说话，系统现在处于聆听、思考还是回答状态，不同智能体之间是否共享了前文信息。"),
        ("p", "由此，本章将多智能体交互设计拆为两个层次：外层是外观、屏幕和角色形象，用于帮助用户识别不同智能体；内层是对话状态、信息共享和上下文边界，用于帮助用户理解智能体之间的关系。前者解决“看得见谁在服务”，后者解决“理解系统如何协作”。"),
        ("h2", "4.2 外观与交互形象研究"),
        ("p", "外观研究关注的是，多智能体语音设备应当以什么样的形态进入家庭空间。对于单智能体音箱而言，无屏或极简灯效通常足以表达设备是否在工作；但对于多智能体音箱而言，用户还需要区分不同角色的身份、人格和服务边界。因此，本研究以“显示能力”和“智能体人格具象程度”为两个维度，探索不同音箱外观方案。"),
        ("p", "显示能力维度包括无屏、小屏、圆形屏和更强显示能力等方向；人格具象程度维度则从抽象状态提示到具体头像、动画形象不等。外观研究关注的问题包括：屏幕是否有助于用户判断当前智能体身份；头像和动画是否提升角色存在感；过度拟人化是否会造成不适；不同显示能力是否会改变用户对设备专业性、亲和力和家居适配性的评价。"),
        ("p", "该研究得到的结论是：多智能体音箱需要一定程度的视觉表达，但不一定需要大屏或过度拟人化外观。小型屏幕配合稳定的智能体头像、名称和状态动效，能够在不破坏家居设备日常感的前提下，帮助用户快速判断当前角色。这一结论直接影响后续原型选择，即采用带 2.8 寸圆形 IPS 屏幕的音箱方案来呈现智能体身份与状态。"),
        ("h2", "4.3 屏幕显示与信息共享实验"),
        ("p", "在外观研究之后，本研究进一步通过原型实验考察屏幕显示和信息共享对多智能体语音交互体验的影响。该实验并不是单纯验证“有没有屏幕更好”，而是进一步考察屏幕是否能帮助用户理解多智能体切换，以及智能体之间共享上下文是否会在提升对话流畅度的同时带来隐私和边界感担忧。"),
        ("h3", "4.3.1 实验目标与总体设计"),
        ("p", "实验采用 2×2 被试内设计，自变量包括外部视觉显示模式和智能体间信息传输模式。外部视觉显示分为有屏幕和无屏幕两种条件：有屏幕条件下，系统会显示当前智能体头像、名称和状态；无屏幕条件下，用户只能通过语音判断当前交互状态。智能体间信息传输模式分为共享和隔离两种条件：共享条件下，主智能体可以获取用户与子智能体之间的前序对话，并在再次被唤起时进行上下文关联回应；隔离条件下，主智能体无法访问用户与其他子智能体的对话内容。"),
        ("p", "由此形成四种实验条件，所有参与者均体验四种条件。实验重点观察用户在智能体切换过程中的系统可用性、认知负荷、信任感、隐私感知和识别效率。通过这一设计，可以同时验证屏幕显示是否提升多智能体可理解性，以及信息共享是否在流畅体验和隐私边界之间产生张力。"),
        ("h3", "4.3.2 实验刺激物与智能体设置"),
        ("p", "实验刺激物为 AgentsNest 多智能体语音原型。软件层面，系统基于 Coze 多智能体平台构建，支持实时语音对话、智能体切换、个性化人设回复、语音输出和跨智能体上下文记忆。硬件层面，实验使用智能音箱、IPS 液晶屏、实体按键和 ESP32 控制模块组合搭建，其中屏幕仅在有屏幕条件下启用，用于呈现智能体头像和状态反馈；实体按键用于记录用户识别智能体切换的反应时间。"),
        ("p", "实验系统包含一个主智能体和三个子智能体。主智能体承担通用对话和上下文协调功能，子智能体分别承担不同人设和任务类型，例如孙悟空、占卜师和英语教师。这样的角色组合覆盖娱乐文化、情感表达和学习辅导等不同交互场景，有助于模拟用户在真实多智能体设备中频繁切换角色的体验。"),
        ("h3", "4.3.3 实验流程与测量指标"),
        ("p", "实验流程包括准备、休息说明、对话任务、问卷填写、半结构化访谈和结束六个阶段。参与者首先完成知情同意和基本信息填写，随后依次完成四组实验条件。每组任务中，参与者先与主智能体对话，再依次切换到不同子智能体，最后回到主智能体并询问与前序对话相关的问题。该流程用于检验共享或隔离条件下，主智能体能否利用前序上下文，以及用户如何感知这种信息传输方式。"),
        ("p", "实验从三个维度进行测量。第一是视觉显示维度，主要通过 NASA-TLX 量表测量认知负荷，并通过 Trust in Automation Scale 测量用户对系统的信任感。第二是智能体间信息传输维度，主要通过隐私关注量表测量用户对信息共享、控制感和信息风险的感知。第三是整体交互维度，使用 SUS 测量系统可用性和满意度，同时通过智能体切换识别任务中的反应时间和准确率测量识别效率。"),
        ("h2", "4.4 设计结论与原则"),
        ("p", "定量结果显示，屏幕的存在能够提升系统易用性，减轻用户认知负荷，提高识别效率，并增加用户信任感。说明对于多智能体语音设备而言，屏幕并不只是装饰性元素，而是帮助用户理解系统状态的重要媒介。"),
        ("p", "有旁听的信息共享模式能够提升系统易用性，但它也会降低部分用户的隐私感知。该结果说明，多智能体系统不能简单地追求信息完全共享，而需要在流畅性和用户可控之间取得平衡。用户一方面希望系统像“中心化管家”一样自动协调多个智能体，减少重复说明；另一方面也希望知道哪些智能体在旁听，哪些信息被共享，以及自己是否可以关闭这种共享。"),
        ("p", "结合访谈结果，本研究将多智能体交互设计原则概括为四点。第一，角色可分辨：不同智能体应具有稳定的名称、声音、头像和能力标签。第二，状态可感知：界面应表达当前系统处于聆听、思考、说话、等待还是切换状态。第三，切换可解释：系统切换智能体时，应让用户知道为什么切换，以及是否仍保留前文上下文。第四，边界可控制：用户应能够理解并管理智能体之间的信息共享、旁听和长期记忆权限。"),
        ("p", "因此，第4章的交互研究与第3章的匹配方法形成互补关系：第3章解决系统如何选对智能体，第4章解决用户如何看懂、听懂并控制这个选择。二者共同构成本研究面向需求多样性的多智能体语音体验设计基础。"),
    ]
    replace_range(doc, ch4, ch5, ch4_blocks)

    for para in doc.paragraphs:
        if para.text.strip().startswith("本研究的方法贡献在于提出从即时情境响应、用户偏好、长期记忆到持续个性化"):
            set_para_text(para, "本研究的方法贡献在于提出即时情境响应、用户偏好和长期记忆三层个性化匹配方法。该方法将多模态用户特征识别、大模型语义判断、用户研究量表数据、智能体标签和分层记忆机制结合起来，并通过实验设计验证其相较唤醒词方式是否能降低选择负担、提升匹配合适度。")

    doc.save(str(DOCX))
    print(DOCX)
    print(IMG)


if __name__ == "__main__":
    main()
