from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches


DOCX = Path("/Users/renke/毕设/中期答辩/任科中期报告_0629_新框架版.docx")
IMG = Path("/Users/renke/Documents/Codex/2026-05-27/1-1-2-agent-1-agent/tmp/research_framework_0630.png")
TITLE = "面向需求多样性的AI语音智能体个性化匹配与交互设计研究"


def make_framework_image() -> None:
    from PIL import Image, ImageDraw, ImageFont

    IMG.parent.mkdir(parents=True, exist_ok=True)
    W, H = 1700, 2250
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    font_path = "/System/Library/Fonts/Hiragino Sans GB.ttc"
    title = ImageFont.truetype(font_path, 44)
    h1 = ImageFont.truetype(font_path, 31)
    body = ImageFont.truetype(font_path, 22)
    small = ImageFont.truetype(font_path, 18)
    tiny = ImageFont.truetype(font_path, 16)

    def wrap(text: str, font, max_w: int) -> list[str]:
        lines, cur = [], ""
        for ch in text:
            if ch == "\n":
                lines.append(cur)
                cur = ""
                continue
            test = cur + ch
            if d.textbbox((0, 0), test, font=font)[2] <= max_w:
                cur = test
            else:
                if cur:
                    lines.append(cur)
                cur = ch
        if cur:
            lines.append(cur)
        return lines

    def centered_text(x, y, w, text, font, fill="#222222", line_gap=5):
        lines = wrap(text, font, int(w))
        total = len(lines) * font.size + (len(lines) - 1) * line_gap
        yy = y - total / 2
        for line in lines:
            box = d.textbbox((0, 0), line, font=font)
            d.text((x + (w - (box[2] - box[0])) / 2, yy), line, font=font, fill=fill)
            yy += font.size + line_gap

    def rect(x, y, w, h, fill="#ffffff", outline="#666666", width=2):
        d.rectangle((x, y, x + w, y + h), fill=fill, outline=outline, width=width)

    def line(x1, y1, x2, y2, width=3):
        d.line((x1, y1, x2, y2), fill="#555555", width=width)

    def arrow_down(x, y1, y2):
        line(x, y1, x, y2, 3)
        d.polygon([(x, y2 + 13), (x - 12, y2 - 6), (x + 12, y2 - 6)], fill="#555555")

    def box(text, x, y, w, h, fill="#ffffff", font=None):
        rect(x, y, w, h, fill, "#666666", 2)
        centered_text(x + 12, y + h / 2, w - 24, text, font or body)

    def vlabel(text, x, y, h):
        rect(x, y, 68, h, "#d9d9d9", "#666666", 2)
        yy = y + (h - len(text) * 33) / 2
        for ch in text:
            centered_text(x + 8, yy + 16, 52, ch, h1, fill="#222222", line_gap=0)
            yy += 33

    d.text((110, 54), f"研究框架：{TITLE}", font=title, fill="#111111")
    d.line((110, 124, 1590, 124), fill="#111111", width=6)
    main_w = 1180
    main_x = (W - main_w) // 2
    side_x = 88
    cx = W // 2
    y = 164
    box(TITLE, main_x + 55, y, main_w - 110, 66, "#d9d9d9", h1)
    arrow_down(cx, y + 66, y + 108)
    y += 118

    sections = [
        ("研究问题", "一、研究背景与研究问题", ["需求多样性", "单一智能体不足", "多智能体方案", "交互需优化"], "RQ1 动态匹配合适智能体    RQ2 语音交互可理解"),
        ("体验框架", "二、多智能体个性化体验设计框架", ["用户研究\n访谈+量表", "体验重构\n需求识别→匹配", "个性化匹配设计框架"], "即时情境响应 → 用户偏好 → 长期记忆 → 持续个性化"),
        ("匹配方法", "三、个性化匹配方法构建", ["Context Matching\n多模态识别+语义", "Preference Matching\n量表偏好数据", "Memory Matching\npersona/context/上下文", "Long-term Personalization\n持续学习"], "从当下情境判断，走向可长期演进的个性化匹配"),
        ("交互设计", "四、多智能体交互设计研究", ["交互形象\n身份可分辨", "模态形式\n声音+屏幕+动效", "状态反馈\n听/想/说", "切换边界\n保持/切换"], "让用户理解当前是谁、为何切换、何时保持对话"),
        ("原型设计", "五、原型设计", ["系统架构", "音箱外观", "音箱UI", "手机端APP"], "形成可测试的带屏多智能体语音音箱与控制台"),
        ("实验验证", "六、实验验证", ["原型体验验证\n新原型 vs 小米", "方法有效性验证\n个性化匹配 vs 唤醒词", "访谈\n理解与信任"], "验证体验价值、匹配有效性与交互可理解性"),
        ("研究产出", "七、研究产出总结", ["理论贡献", "方法贡献", "设计贡献", "实践贡献"], "沉淀多智能体语音交互体验框架、方法与原则"),
    ]

    for idx, (side, title_text, nodes, note) in enumerate(sections):
        sec_h = 240 if idx in (0, 2, 3) else 224
        sec_y = y
        vlabel(side, side_x, sec_y + 28, sec_h - 56)
        rect(main_x, sec_y, main_w, sec_h)
        box(title_text, main_x + 340, sec_y + 24, 500, 54, "#d9d9d9", h1)
        n = len(nodes)
        gap = 30
        bw = int((main_w - 140 - gap * (n - 1)) / n)
        start_x = main_x + 70
        line_y = sec_y + 108
        line(start_x + bw / 2, line_y, start_x + (n - 1) * (bw + gap) + bw / 2, line_y, 3)
        line(cx, sec_y + 78, cx, line_y, 3)
        for i, text in enumerate(nodes):
            x = start_x + i * (bw + gap)
            box(text, x, sec_y + 128, bw, 66, "#ffffff", small)
            line(x + bw / 2, line_y, x + bw / 2, sec_y + 128, 2)
        box(note, main_x + 250, sec_y + sec_h - 38, 680, 32, "#f2f2f2", tiny)
        if idx < len(sections) - 1:
            arrow_down(cx, sec_y + sec_h, sec_y + sec_h + 38)
        y += sec_h + 54
    img.save(IMG)


def set_para_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_paragraph_before(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = paragraph.__class__(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def add_before(anchor, typ: str, text: str = ""):
    style = {
        "h1": "Report Heading 1",
        "h2": "Report Heading 2",
        "p": "Report Body",
        "toc": "Report TOC",
    }[typ]
    return insert_paragraph_before(anchor, text, style)


def replace_range_before_anchor(doc: Document, start_idx: int, end_idx: int, blocks: list[tuple[str, str]]) -> None:
    anchor = doc.paragraphs[end_idx]
    for para in doc.paragraphs[start_idx:end_idx]:
        para._element.getparent().remove(para._element)
    for typ, text in blocks:
        add_before(anchor, typ, text)


def replace_picture(paragraph, path: Path) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.65))


def update_footer(doc: Document, total: str = "19") -> None:
    for section in doc.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for para in footer.paragraphs:
                for run in para.runs:
                    if "39" in run.text:
                        run.text = run.text.replace("39", total)


def main() -> None:
    make_framework_image()
    doc = Document(str(DOCX))

    # TOC entries only: keep manual formatting.
    toc_replacements = {
        "2.3 个性化匹配演进框架": "2.3 个性化匹配设计框架",
        "3.1 情境化匹配": "3.1 即时的情境响应（Context Matching）",
        "3.2 基于长期记忆匹配": "3.2 用户偏好（Preference Matching）\n3.3 长期记忆（Memory Matching）\n3.4 持续个性化（Long-term Personalization）",
        "第4章 多智能体交互设计研究": "第4章 多智能体交互设计研究\n4.1 多智能体交互形象设计\n4.2 多模态状态反馈设计\n4.3 会话保持与智能体切换设计\n4.4 信息边界与用户控制",
    }
    for para in doc.paragraphs:
        if para.style.name == "Report TOC":
            text = para.text.strip()
            if text in toc_replacements:
                set_para_text(para, toc_replacements[text])

    for i, para in enumerate(doc.paragraphs):
        t = para.text.strip()
        if t == "2.3 个性化匹配演进框架":
            set_para_text(para, "2.3 个性化匹配设计框架")
        elif t.startswith("为了让个性化匹配具有可持续扩展性"):
            set_para_text(para, "为了让个性化匹配具有可持续扩展性，本研究将其表述为“个性化匹配设计框架”。该框架强调从情境化匹配到基于长期记忆匹配的体验演进路线，分为：即时的情境响应（Context Matching）、用户偏好（Preference Matching）、长期记忆（Memory Matching）和持续个性化（Long-term Personalization）。")
        elif t.startswith("即时的情境响应（Context Matching） 指"):
            set_para_text(para, "即时的情境响应（Context Matching）指系统根据当前一次语音输入、当前时间、用户临时状态和对话语境进行即时判断；用户偏好（Preference Matching）指系统结合前序访谈与量表数据，根据用户类型和场景对智能体特征进行偏好排序；长期记忆（Memory Matching）指系统在用户授权下持续维护 persona、近期话题和临时上下文，使匹配不只依赖当下 query；持续个性化（Long-term Personalization）则指系统根据长期反馈不断调整匹配策略、默认智能体和提醒方式，使家庭语音系统逐渐贴合具体家庭的使用习惯。")
        elif t == "围绕这一重构，本研究同步设计了用户旅程、语音交互流程。":
            set_para_text(para, "围绕这一重构，本研究同步设计了用户旅程、语音交互流程和多智能体语音对话架构。")
        elif t.startswith("语音交互流程用于明确系统"):
            set_para_text(para, "语音交互流程用于明确系统在强意图、弱意图、会话保持和异步任务中的状态变化；语音对话架构则用于梳理用户如何在与多个智能体的交互过程中保持会话流畅，以及如何在智能体切换、会话保持、打断恢复和异步提醒中获得无缝的对话体验。")

    # Replace the framework image paragraph in its current position.
    for i, para in enumerate(doc.paragraphs[:-1]):
        if para._p.xml.count("graphic") > 0 and doc.paragraphs[i + 1].text.strip() == "图1 研究框架图":
            replace_picture(para, IMG)
            break

    # Recompute indices after previous edits.
    idx = {p.text.strip(): i for i, p in enumerate(doc.paragraphs) if p.text.strip()}

    ch3 = idx["第3章 个性化匹配方法构建"]
    ch4 = idx["第4章 多智能体交互设计研究"]
    ch5 = idx["第5章 原型设计"]

    ch3_blocks = [
        ("h1", "第3章 个性化匹配方法构建"),
        ("p", "在体验框架的基础上，本研究将个性化匹配方法进一步拆分为四个层次：即时的情境响应、用户偏好、长期记忆和持续个性化。四个层次并不是彼此替代，而是从“当前这句话应该由谁回答”逐步扩展到“系统如何随着长期使用逐渐更懂这个家庭”。"),
        ("h2", "3.1 即时的情境响应（Context Matching）"),
        ("p", "即时的情境响应主要解决用户当下说出一句话时，系统如何快速判断合适的智能体。该层重点依据当前输入，而不是长期偏好。首先，系统通过语音或多模态方式识别用户的粗粒度特征，例如年龄段和性别，用于判断儿童、成人或中老年用户在同一 query 下可能需要的不同回应方式。其次，系统结合当前时间判断早、中、晚等使用情境，因为同样一句“陪我一下”在睡前、午后和早晨可能对应不同的陪伴方式。"),
        ("p", "在语义层面，系统使用大模型对用户 query 进行语义识别与匹配，判断其更接近休闲陪伴、健康咨询、智能家居、信息查询或学习辅导等场景。这里不只依赖关键词，而是关注用户真实需求。例如“我现在需要陪小孩”虽然没有直接说“讲故事”，但在儿童陪伴语境中更可能需要儿童友好的智能体，而不是健康顾问或通用助手。"),
        ("p", "即时情境响应还需要区分强意图和弱意图。强意图是用户明确说出智能体名称、别名或唤醒词，这类请求直接路由；弱意图是用户自然表达需求但没有指定智能体，需要系统结合多模态特征、当前时间和语义判断完成初步匹配。"),
        ("h2", "3.2 用户偏好（Preference Matching）"),
        ("p", "用户偏好匹配用于把前序用户研究和量表数据真正转化为系统行为。本研究在量表中将智能体特征归纳为亲密陪伴型、鼓励赞赏型、成熟关怀型、专业可靠型和中性工具型，并获得不同用户类型、不同场景下对这些特征的偏好分。系统在弱意图匹配时，会先根据用户类型和场景查找偏好分，再结合当前智能体池中每个智能体的特征标签进行排序。"),
        ("p", "具体而言，每个智能体可以被标注为一种或两种偏好类型。例如儿童陪伴智能体可能对应亲密陪伴型和鼓励赞赏型，健康顾问可能对应专业可靠型和成熟关怀型。系统根据用户类型与场景得到各类智能体特征的偏好分，再对智能体标签取平均或加权计算，得到候选智能体排序。该步骤通常召回 Top3 候选，再交给大模型结合 query、上下文和智能体描述做最终选择。"),
        ("p", "这一层的意义在于避免模型完全凭语义直觉选择智能体，而是引入用户研究得到的先验偏好。它使匹配逻辑能够解释为：在某类用户、某类场景中，用户更偏好什么样的智能体特征，因此系统优先推荐与这些特征更匹配的智能体。"),
        ("h2", "3.3 长期记忆（Memory Matching）"),
        ("p", "长期记忆匹配用于解决一次性情境判断无法覆盖的长期家庭使用问题。本研究将记忆分为三类：persona、context 和临时上下文。三者的更新规则和频率不同，因此不能混为一种记忆。"),
        ("p", "persona 指相对稳定的用户画像，包括家庭成员身份、年龄段、常用智能体、偏好的语气、拒绝的风格、常见场景和隐私边界。这类信息更新频率较低，应当通过明确反馈、长期统计或用户在手机端设置后更新，避免系统因为一次偶然对话就改变用户画像。context 指最近一段时间内的对话主题和任务状态，例如最近在聊英语学习、正在进行健康咨询、刚刚设置了看电视提醒等。这类信息更新频率中等，随着会话结束、任务完成或时间过去而衰减。临时上下文则指当前一轮对话中的短时信息，例如上一句话、当前智能体、是否处于20秒会话保持期、用户是否刚刚打断。这类信息更新频率最高，通常只在当前会话窗口内有效。"),
        ("p", "通过区分三类记忆，系统可以在不牺牲隐私和稳定性的前提下提升匹配连续性。例如，用户短时间内继续提问时，系统优先参考临时上下文保持当前智能体；用户隔天再次提出相似需求时，系统可参考 persona 和长期偏好；任务提醒到点时，系统参考 context 中的任务状态和提醒对象，选择更合适的智能体进行提醒。"),
        ("h2", "3.4 持续个性化（Long-term Personalization）"),
        ("p", "持续个性化关注的是系统如何在长期使用中不断调整匹配策略。它与长期记忆的区别在于：长期记忆更关注“记住什么”，持续个性化更关注“如何根据记忆改变后续系统行为”。例如，系统记住某个儿童更喜欢虎子讲故事属于长期记忆；之后在儿童休闲陪伴场景中默认优先推荐虎子，并减少重复询问，则属于持续个性化。"),
        ("p", "持续个性化需要建立反馈闭环。用户可以通过继续使用、打断、手动切换、拒绝推荐或在手机端调整偏好等方式向系统提供反馈。系统根据这些反馈逐步调整默认智能体、候选排序、提醒方式和信息共享边界。这样，多智能体音箱不只是一次性完成路由，而是在家庭长期使用中逐渐形成更稳定、更自然的个性化服务。"),
    ]
    replace_range_before_anchor(doc, ch3, ch4, ch3_blocks)

    # Recompute for chapter 4 after replacing chapter 3.
    idx = {p.text.strip(): i for i, p in enumerate(doc.paragraphs) if p.text.strip()}
    ch4 = idx["第4章 多智能体交互设计研究"]
    ch5 = idx["第5章 原型设计"]
    ch4_blocks = [
        ("h1", "第4章 多智能体交互设计研究"),
        ("p", "多智能体系统如果只在后台完成匹配，用户仍然会觉得它像黑箱。因此，本研究进一步研究如何重新设计多智能体语音硬件的交互形象与模态形式。这里的核心出发点是当多个角色共同存在时，用户需要什么样的交互形象和信息表达，才能理解当前系统。"),
        ("h2", "4.1 多智能体交互形象设计"),
        ("p", "多智能体交互形象设计关注不同智能体如何被用户区分。前期人机交互实验围绕音箱显示能力和智能体人格具象程度展开。显示能力包括无屏、小屏、圆形屏和更强显示能力等方向；人格具象程度则从抽象状态提示到具体头像、动画形象不等。研究关注的问题包括：屏幕是否有助于用户判断当前智能体身份；头像和动画是否提升角色存在感；过度拟人化是否会造成不适。"),
        ("p", "研究结果支持带屏音箱作为多智能体语音交互的重要载体。屏幕不需要承担复杂操作，但需要清楚传达当前智能体、状态变化和切换反馈。不同智能体应当具有稳定的名称、头像、语音风格和视觉表达，使用户能够形成可持续的角色认知。"),
        ("h2", "4.2 多模态状态反馈设计"),
        ("p", "多智能体语音系统中的状态反馈比单智能体更重要。用户不仅需要知道系统是否听到了自己，还需要知道当前是哪一个智能体在听、在思考、在说话。音箱 UI 因此需要通过头像、动效、颜色、文字和语音提示表达 listening、thinking、speaking、waiting 等状态，减少用户在等待和切换过程中的不确定感。"),
        ("h2", "4.3 会话保持与智能体切换设计"),
        ("p", "多智能体系统不能每句话都重新匹配，否则会造成体验割裂。因此，本研究引入会话保持机制：当某个智能体回复结束后，在短时间内如果用户继续提问且没有明确唤醒其他智能体，则由当前智能体继续回答；只有当用户出现强意图、明显跨场景需求或会话超时，系统才重新匹配。该机制需要在界面上被表达出来，让用户理解什么时候系统会保持当前对话，什么时候会切换到其他智能体。"),
        ("h2", "4.4 信息边界与用户控制"),
        ("p", "当多个智能体共同存在时，用户会关心前文信息是否会被其他智能体知道，以及哪些信息被用于后续匹配。因此，多智能体交互设计还需要表达信息边界。手机端可以承担更完整的隐私与记忆管理功能，例如查看 persona、近期 context、任务记忆和智能体访问权限；音箱端则通过轻量提示说明当前是否发生智能体切换、是否沿用上下文、是否记录长期偏好。"),
        ("p", "因此，本研究将多智能体交互设计原则初步概括为四点：角色可分辨、状态可感知、切换可解释、边界可控制。这些原则共同回应 RQ2，即如何让多智能体语音交互可理解。"),
    ]
    replace_range_before_anchor(doc, ch4, ch5, ch4_blocks)

    for para in doc.paragraphs:
        if para.text.strip() == "本研究的方法贡献在于提出从情境化匹配到长期记忆匹配的个性化匹配方法。该方法将用户研究得到的偏好数据、用户特征、场景识别、智能体标签和大模型语义判断结合起来，并通过实验设计验证其相较唤醒词方式是否能降低选择负担、提升匹配合适度。":
            set_para_text(para, "本研究的方法贡献在于提出从即时情境响应、用户偏好、长期记忆到持续个性化的个性化匹配方法。该方法将多模态用户特征识别、大模型语义判断、用户研究量表数据、智能体标签和分层记忆机制结合起来，并通过实验设计验证其相较唤醒词方式是否能降低选择负担、提升匹配合适度。")

    update_footer(doc, "19")
    doc.save(str(DOCX))
    print(DOCX)
    print(IMG)


if __name__ == "__main__":
    main()
