from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOCX = Path("/Users/renke/毕设/中期答辩/任科中期报告_0629_新框架版.docx")


def insert_after(paragraph: Paragraph, text: str, style: str = "Report Body") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = style
    inserted.add_run(text)
    return inserted


def replace_exact(doc: Document, old: str, new: str) -> int:
    count = 0
    for paragraph in doc.paragraphs:
        if paragraph.text == old:
            paragraph.text = new
            count += 1
    return count


def find_paragraph(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text == text:
            return paragraph
    raise ValueError(f"paragraph not found: {text[:50]}")


def main() -> None:
    doc = Document(str(DOCX))

    placeholder = find_paragraph(doc, "关于实现的技术细节，")
    placeholder.text = (
        "关于实现的技术细节，长期记忆模块可以理解为系统内部的“家庭使用档案”，但它并不是把所有对话原文无限保存，"
        "而是按信息稳定程度进行分层管理。系统为每个家庭成员维护 userId、persona、context、sessionState、"
        "permission 和 updatedAt 等字段：persona 用于记录相对稳定的画像与偏好，context 用于记录近期主题和任务状态，"
        "sessionState 用于记录当前一轮对话中的临时上下文，permission 则用于约束哪些记忆可以被哪些智能体使用。"
    )
    p = insert_after(
        placeholder,
        "在更新规则上，persona 只在用户明确设置、明确纠正或多次重复出现同一偏好时更新，例如用户在手机端设置“儿童提醒用温和语气”，"
        "或多次选择虎子讲故事；context 在一次对话结束、任务创建或任务完成后更新，并随时间衰减；临时上下文则在当前会话中即时更新，"
        "例如当前智能体是谁、用户是否处于会话保持期、上一轮是否刚刚打断。这样的设计避免了系统因为一次偶然表达就改变长期判断，"
        "也避免每一次对话都从零开始。"
    )
    p = insert_after(
        p,
        "在匹配阶段，长期记忆不是直接替代用户当下的需求，而是作为弱意图匹配中的辅助依据。系统会先判断强意图；若用户明确唤醒某个智能体，"
        "仍然直接切换。只有在弱意图或需要自动推荐时，路由模块才会把当前 query、用户画像、近期 context、可用智能体标签和隐私权限组合成"
        "一个 memory-aware routing packet，再交给偏好排序和大模型做最终判断。对于设计学语境而言，可以把它理解为：系统先听懂用户此刻"
        "要做什么，再参考“这个家庭过去通常怎么做”，最后选择一个更合适、更少打扰的智能体。"
    )
    insert_after(
        p,
        "此外，长期记忆还承担异步任务的承接作用。例如家长说“提醒小朋友 20 分钟后写作业”，系统会生成包含 triggerAt、audience、message、"
        "reminderAgentId 和 status 的任务对象；到达时间后，系统不是简单播放闹钟，而是结合 audience 和记忆中的偏好，选择更适合儿童的智能体"
        "进行提醒。为了降低隐私风险，健康咨询、儿童信息和家庭任务等记忆应在手机端提供查看、修改和删除入口，并允许用户关闭跨智能体信息共享。"
    )

    chapter5 = find_paragraph(doc, "第5章 原型设计")
    insert_after(
        chapter5,
        "前文已经分别完成了个性化匹配方法和多智能体交互设计原则的构建。第 3 章回答系统如何根据用户需求、偏好和长期记忆选择合适智能体，"
        "第 4 章回答用户如何看懂当前智能体身份、状态和信息共享边界。因此，本章的原型设计并不是单纯做一个外观或界面展示，而是把前文的方法"
        "和交互原则转化为可体验、可测试的系统载体。"
    )

    replace_exact(
        doc,
        "音箱外观设计围绕“显示能力”和“智能体人格具象程度”两个维度展开。经过多方案比较，本研究选择带有 2.8 寸圆形 IPS 屏幕的智能音箱方向，用于显示当前智能体头像和状态动效。该方案保留了智能音箱作为家居设备的简洁感，同时为多智能体交互提供必要的身份和状态表达空间。",
        "基于第 4 章外观与交互形象研究的结论，音箱外观设计围绕“显示能力”和“智能体人格具象程度”两个维度展开。前一章研究表明，多智能体音箱需要一定程度的视觉表达来帮助用户识别当前角色，但又不宜过度拟人化，以免削弱其作为家庭设备的稳定感和适配性。经过多方案比较，本研究选择带有 2.8 寸圆形 IPS 屏幕的智能音箱方向，用于显示当前智能体头像和状态动效。该方案保留了智能音箱作为家居设备的简洁感，同时为多智能体交互提供必要的身份和状态表达空间。"
    )

    doc.save(str(DOCX))
    print(DOCX)


if __name__ == "__main__":
    main()
