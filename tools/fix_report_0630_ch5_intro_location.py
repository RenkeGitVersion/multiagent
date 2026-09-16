from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOCX = Path("/Users/renke/毕设/中期答辩/任科中期报告_0629_新框架版.docx")
INTRO = (
    "前文已经分别完成了个性化匹配方法和多智能体交互设计原则的构建。第 3 章回答系统如何根据用户需求、偏好和长期记忆选择合适智能体，"
    "第 4 章回答用户如何看懂当前智能体身份、状态和信息共享边界。因此，本章的原型设计并不是单纯做一个外观或界面展示，而是把前文的方法"
    "和交互原则转化为可体验、可测试的系统载体。"
)


def insert_after(paragraph: Paragraph, text: str, style: str = "Report Body") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = style
    inserted.add_run(text)
    return inserted


def delete_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def main() -> None:
    doc = Document(str(DOCX))

    for paragraph in list(doc.paragraphs):
        if paragraph.text == INTRO:
            delete_paragraph(paragraph)

    headings = [p for p in doc.paragraphs if p.text == "第5章 原型设计"]
    target = next(p for p in headings if p.style.name == "Report Heading 1")
    insert_after(target, INTRO)

    doc.save(str(DOCX))
    print(DOCX)


if __name__ == "__main__":
    main()
