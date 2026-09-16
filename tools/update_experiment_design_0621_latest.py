from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SRC = Path("/Users/renke/毕设/中期答辩/任科中期报告_0621_最新.docx")
BACKUP = Path("/Users/renke/毕设/中期答辩/任科中期报告_0621_最新_备份_修改前.docx")


def norm(text: str) -> str:
    return " ".join(text.split())


def set_run_font(run, east="仿宋_GB2312", ascii_font="Times New Roman", size=12, bold=False):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold


def clear_and_write(paragraph, text: str):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run)


def insert_before(paragraph, text: str, style, bold_prefix: str | None = None):
    parent = paragraph._p.getparent()
    idx = parent.index(paragraph._p)
    p_el = OxmlElement("w:p")
    parent.insert(idx, p_el)
    doc = paragraph.part.document
    new_p = next(p for p in doc.paragraphs if p._p is p_el)
    new_p.style = style
    if style.name == "Report Body":
        new_p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = new_p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = new_p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = new_p.add_run(text)
        set_run_font(r, bold=(style.name == "Report Heading 3"))
    return new_p


def main():
    if not BACKUP.exists():
        shutil.copy2(SRC, BACKUP)

    doc = Document(SRC)
    paras = doc.paragraphs
    start = next(i for i, p in enumerate(paras) if norm(p.text) == "2.3.3 实验设计")
    end = next(i for i, p in enumerate(paras[start + 1 :], start=start + 1) if norm(p.text) == "2.3.4 定量结论")

    body_style = doc.styles["Report Body"]

    # Replace the original brief paragraph, and insert expanded content before existing figures/captions.
    first_body = next(p for p in paras[start + 1 : end] if norm(p.text) and p.style.name == "Report Body")
    clear_and_write(
        first_body,
        "本研究参考论文中 AgentsNest 实验设计，采用受控实验方式评估多智能体语音硬件设备中“外部视觉显示”和“智能体间信息传输模式”对用户体验的影响。实验以多智能体语音交互为任务情境，重点观察用户在智能体切换过程中的系统可用性、认知负荷、信任感、隐私感知和识别效率。",
    )

    anchor = next(p for p in paras[start + 1 : end] if p.style.name == "Report Caption" or not norm(p.text))

    additions = [
        ("实验变量：实验采用 2×2 被试内设计，自变量包括外部视觉显示模式和智能体间信息传输模式。外部视觉显示分为有屏幕和无屏幕两种条件：有屏幕条件下，系统会显示当前智能体头像、名称和状态；无屏幕条件下，用户只能通过语音判断当前交互状态。智能体间信息传输模式分为共享和隔离两种条件：共享条件下，主智能体可以获取用户与子智能体之间的前序对话，并在再次被唤起时进行上下文关联回应；隔离条件下，主智能体无法访问用户与其他子智能体的对话内容。由此形成四种实验条件，所有参与者均体验四种条件。", "实验变量："),
        ("被试与实验顺序：实验采用被试内设计，以减少个体差异对结果的影响。根据 G*Power 3.1.9.7 对重复测量方差分析的样本量估计，实验计划招募 28 名参与者。正式实验中，参与者年龄集中在 20-25 岁，均具备智能语音助手使用经验。为控制顺序效应，四种实验条件通过拉丁方进行平衡，使不同参与者以不同顺序体验各条件。", "被试与实验顺序："),
        ("实验刺激物：实验刺激物为 AgentsNest 多智能体语音原型。软件层面，系统基于 Coze 多智能体平台构建，能够支持实时语音对话、智能体切换、个性化人设回复、语音输出和跨智能体上下文记忆。硬件层面，实验使用智能音箱、IPS 液晶屏、实体按键和 ESP32 控制模块组合搭建，其中屏幕仅在有屏幕条件下启用，用于呈现智能体头像和状态反馈；实体按键用于记录用户识别智能体切换的反应时间。", "实验刺激物："),
        ("智能体设置：实验系统包含一个主智能体和三个子智能体。主智能体承担通用对话和上下文协调功能，子智能体分别承担不同人设和任务类型，例如孙悟空、占卜师和英语教师。这样的角色组合覆盖娱乐文化、情感表达和学习辅导等不同交互场景，有助于模拟用户在真实多智能体设备中频繁切换角色的体验。", "智能体设置："),
        ("实验流程：实验流程包括准备、休息说明、对话任务、问卷填写、半结构化访谈和结束六个阶段。参与者首先完成知情同意和基本信息填写，随后依次完成四组实验条件。每组任务中，参与者先与主智能体对话，再依次切换到不同子智能体，最后回到主智能体并询问与前序对话相关的问题。该流程用于检验共享或隔离条件下，主智能体能否利用前序上下文，以及用户如何感知这种信息传输方式。", "实验流程："),
        ("实验控制：为保证不同被试和不同条件下的交互内容、节奏和视觉反馈具有可比性，实验采用 Wizard-of-Oz 控制方式。实验员根据预设脚本选择智能体回复，手动控制屏幕反馈和动画切换，并在每轮实验前清除对话记忆，以模拟不同信息传输条件。这样既保留了自然语音交互体验，又能控制实验变量，避免由模型生成内容差异造成额外干扰。", "实验控制："),
        ("因变量与量表：实验从三个维度进行测量。第一是视觉显示维度，主要通过 NASA-TLX 量表测量认知负荷，并通过 Trust in Automation Scale 测量用户对系统的信任感。第二是智能体间信息传输维度，主要通过隐私关注量表测量用户对信息共享、控制感和信息风险的感知。第三是整体交互维度，使用 System Usability Scale（SUS）测量系统可用性和满意度，同时通过智能体切换识别任务中的反应时间和准确率测量识别效率。除行为指标外，主观量表均采用 5 点李克特量表。", "因变量与量表："),
        ("访谈内容：完成四种实验条件后，参与者接受半结构化访谈。访谈重点包括用户是否能理解多智能体的身份与状态、是否认为屏幕反馈有助于降低切换负担、是否接受主智能体旁听或共享上下文，以及在未来产品中希望如何控制不同智能体之间的信息传递。", "访谈内容："),
    ]

    for text, prefix in additions:
        insert_before(anchor, text, body_style, prefix)

    doc.save(SRC)
    print(SRC)
    print(BACKUP)


if __name__ == "__main__":
    main()
