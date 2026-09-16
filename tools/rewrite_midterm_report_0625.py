from __future__ import annotations

import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches


SRC = Path("/Users/renke/毕设/中期答辩/任科中期报告_0623.docx")
OUT = Path("/Users/renke/毕设/中期答辩/任科中期报告_0625_重写版.docx")
IMG = Path("/Users/renke/Documents/Codex/2026-05-27/1-1-2-agent-1-agent/tmp/research_framework_0625.png")


def make_framework_image() -> None:
    from PIL import Image, ImageDraw, ImageFont

    IMG.parent.mkdir(parents=True, exist_ok=True)
    W, H = 1700, 2350
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)

    font_path = "/System/Library/Fonts/Hiragino Sans GB.ttc"
    title = ImageFont.truetype(font_path, 46)
    h1 = ImageFont.truetype(font_path, 32)
    body = ImageFont.truetype(font_path, 24)
    small = ImageFont.truetype(font_path, 20)

    def wrap(text: str, font, max_w: int) -> list[str]:
        lines, cur = [], ""
        for ch in text:
            if ch == "\n":
                lines.append(cur)
                cur = ""
                continue
            test = cur + ch
            if d.textbbox((0, 0), test, font=font)[2] <= max_w:
                cur = test
            else:
                if cur:
                    lines.append(cur)
                cur = ch
        if cur:
            lines.append(cur)
        return lines

    def centered_text(x, y, w, text, font, fill=(20, 28, 38), line_gap=8):
        lines = wrap(text, font, int(w))
        total = len(lines) * font.size + (len(lines) - 1) * line_gap
        yy = y - total / 2
        for line in lines:
            box = d.textbbox((0, 0), line, font=font)
            d.text((x + (w - (box[2] - box[0])) / 2, yy), line, font=font, fill=fill)
            yy += font.size + line_gap

    def rect(x, y, w, h, fill="#ffffff", outline="#666666", width=2):
        d.rectangle((x, y, x + w, y + h), fill=fill, outline=outline, width=width)

    def line(x1, y1, x2, y2, width=3):
        d.line((x1, y1, x2, y2), fill="#555555", width=width)

    def arrow_down(x, y1, y2):
        line(x, y1, x, y2, 3)
        d.polygon([(x, y2 + 14), (x - 12, y2 - 6), (x + 12, y2 - 6)], fill="#555555")

    def vlabel(text, x, y, h):
        rect(x, y, 68, h, "#d9d9d9", "#666666", 2)
        chars = list(text)
        total = len(chars) * 34
        yy = y + (h - total) / 2
        for ch in chars:
            centered_text(x + 8, yy + 17, 52, ch, h1, fill="#222222", line_gap=0)
            yy += 34

    def box(text, x, y, w, h, fill="#ffffff", font=None, line_gap=7):
        rect(x, y, w, h, fill, "#666666", 2)
        centered_text(x + 16, y + h / 2, w - 32, text, font or body, fill="#222222", line_gap=line_gap)

    d.text((110, 54), "研究框架：面向需求多样性的AI语音智能体个性化匹配与交互设计研究", font=title, fill="#111111")
    d.line((110, 126, 1590, 126), fill="#111111", width=6)

    cx = W // 2
    main_w = 1180
    main_x = cx - main_w // 2
    side_x = 88
    y = 170

    box("面向需求多样性的AI语音智能体个性化匹配与交互设计研究", main_x + 70, y, main_w - 140, 72, "#d9d9d9", h1)
    arrow_down(cx, y + 72, y + 120)
    y += 130

    # 1. 提出问题
    sec_y, sec_h = y, 260
    vlabel("提出问题", side_x, sec_y + 28, 190)
    rect(main_x, sec_y, main_w, sec_h, "#ffffff", "#666666", 2)
    box("一、提出问题", main_x + 410, sec_y + 28, 360, 58, "#ffffff", h1)
    nodes = [
        ("家庭多用户、多场景需求", main_x + 70),
        ("单一智能体适配不足", main_x + 420),
        ("多个智能体带来服务发现\n和角色理解问题", main_x + 770),
    ]
    for text, x in nodes:
        box(text, x, sec_y + 128, 300, 74, "#ffffff", body)
        line(x + 150, sec_y + 86, x + 150, sec_y + 128)
    line(main_x + 220, sec_y + 104, main_x + 920, sec_y + 104)
    arrow_down(cx, sec_y + sec_h, sec_y + sec_h + 42)
    y += sec_h + 58

    # 2. 研究问题
    sec_y, sec_h = y, 230
    vlabel("研究问题", side_x, sec_y + 24, 180)
    rect(main_x, sec_y, main_w, sec_h, "#ffffff", "#666666", 2)
    box("二、研究问题", main_x + 410, sec_y + 24, 360, 58, "#ffffff", h1)
    box("RQ1：如何匹配合适智能体？", main_x + 190, sec_y + 128, 360, 64, "#ffffff", body)
    box("RQ2：如何让多智能体交互可理解？", main_x + 630, sec_y + 128, 400, 64, "#ffffff", body)
    line(cx, sec_y + 82, cx, sec_y + 112)
    line(main_x + 370, sec_y + 112, main_x + 830, sec_y + 112)
    line(main_x + 370, sec_y + 112, main_x + 370, sec_y + 128)
    line(main_x + 830, sec_y + 112, main_x + 830, sec_y + 128)
    arrow_down(cx, sec_y + sec_h, sec_y + sec_h + 42)
    y += sec_h + 58

    # 3. 设计依据研究
    sec_y, sec_h = y, 285
    vlabel("设计依据", side_x, sec_y + 44, 200)
    rect(main_x, sec_y, main_w, sec_h, "#ffffff", "#666666", 2)
    box("三、设计依据研究", main_x + 380, sec_y + 24, 420, 58, "#d9d9d9", h1)
    basis = [
        ("竞品与案例分析\n现有产品差异", main_x + 80),
        ("用户研究\n场景与偏好", main_x + 400),
        ("人机交互研究\n角色、状态、边界", main_x + 720),
    ]
    for text, x in basis:
        box(text, x, sec_y + 128, 280, 82, "#ffffff", body)
        line(x + 140, sec_y + 82, x + 140, sec_y + 128)
    line(main_x + 220, sec_y + 100, main_x + 860, sec_y + 100)
    box("形成设计依据：服务差异、用户偏好、交互可理解性", main_x + 260, sec_y + 232, 660, 42, "#f2f2f2", small)
    arrow_down(cx, sec_y + sec_h, sec_y + sec_h + 42)
    y += sec_h + 58

    # 4. 方法构建
    sec_y, sec_h = y, 285
    vlabel("方法构建", side_x, sec_y + 44, 200)
    rect(main_x, sec_y, main_w, sec_h, "#ffffff", "#666666", 2)
    box("四、方法构建", main_x + 410, sec_y + 24, 360, 58, "#d9d9d9", h1)
    methods = [
        ("智能体角色池", main_x + 90),
        ("个性化匹配机制", main_x + 360),
        ("多智能体交互表达方法", main_x + 660),
        ("技术原理支撑", main_x + 930),
    ]
    for text, x in methods:
        box(text, x, sec_y + 118, 230, 68, "#ffffff", body)
        line(x + 115, sec_y + 82, x + 115, sec_y + 118)
    line(main_x + 205, sec_y + 96, main_x + 1045, sec_y + 96)
    box("从“用户找服务”转向“系统按需求匹配服务”", main_x + 320, sec_y + 218, 540, 46, "#f2f2f2", small)
    arrow_down(cx, sec_y + sec_h, sec_y + sec_h + 42)
    y += sec_h + 58

    # 5. 原型设计
    sec_y, sec_h = y, 260
    vlabel("原型设计", side_x, sec_y + 34, 190)
    rect(main_x, sec_y, main_w, sec_h, "#ffffff", "#666666", 2)
    box("五、原型设计", main_x + 410, sec_y + 24, 360, 58, "#d9d9d9", h1)
    proto = [
        ("带屏智能音箱", main_x + 120),
        ("手机端控制台", main_x + 470),
        ("多智能体语音系统", main_x + 820),
    ]
    for text, x in proto:
        box(text, x, sec_y + 132, 260, 68, "#ffffff", body)
        line(x + 130, sec_y + 82, x + 130, sec_y + 132)
    line(main_x + 250, sec_y + 102, main_x + 950, sec_y + 102)
    arrow_down(cx, sec_y + sec_h, sec_y + sec_h + 42)
    y += sec_h + 58

    # 6. 实验验证
    sec_y, sec_h = y, 260
    vlabel("实验验证", side_x, sec_y + 34, 190)
    rect(main_x, sec_y, main_w, sec_h, "#ffffff", "#666666", 2)
    box("六、实验验证", main_x + 410, sec_y + 24, 360, 58, "#d9d9d9", h1)
    exp = [
        ("可用性验证\n普通小米音箱 vs 本研究原型", main_x + 90),
        ("匹配有效性验证\n个性化匹配 vs 手动切换", main_x + 420),
        ("用户访谈\n理解原因与体验反馈", main_x + 750),
    ]
    for text, x in exp:
        box(text, x, sec_y + 126, 300, 78, "#ffffff", small)
        line(x + 150, sec_y + 82, x + 150, sec_y + 126)
    line(main_x + 240, sec_y + 100, main_x + 900, sec_y + 100)
    arrow_down(cx, sec_y + sec_h, sec_y + sec_h + 42)
    y += sec_h + 58

    # 7. 研究产出
    sec_y, sec_h = y, 230
    vlabel("研究产出", side_x, sec_y + 24, 180)
    rect(main_x, sec_y, main_w, sec_h, "#ffffff", "#666666", 2)
    box("七、研究产出", main_x + 410, sec_y + 24, 360, 58, "#d9d9d9", h1)
    outputs = [
        ("匹配流程", main_x + 190),
        ("交互原型", main_x + 470),
        ("设计原则", main_x + 750),
    ]
    for text, x in outputs:
        box(text, x, sec_y + 128, 220, 58, "#ffffff", body)
        line(x + 110, sec_y + 82, x + 110, sec_y + 128)
    line(main_x + 300, sec_y + 104, main_x + 860, sec_y + 104)

    img.save(IMG)


def insert_paragraph_before(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = paragraph.__class__(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def clone_run_format(src_run, dst_run):
    if src_run is None or src_run._r.rPr is None:
        return
    dst_run._r.get_or_add_rPr()
    for child in list(dst_run._r.rPr):
        dst_run._r.rPr.remove(child)
    for child in src_run._r.rPr:
        dst_run._r.rPr.append(deepcopy(child))


def add_para(anchor, text="", style="Report Body", align=None):
    p = insert_paragraph_before(anchor, "", style)
    if text:
        run = p.add_run(text)
        # keep direct formatting close to the style defaults
    if align is not None:
        p.alignment = align
    return p


def add_image(anchor, path: Path, width_inches=6.2):
    p = insert_paragraph_before(anchor, "", "Report Body")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width_inches))
    return p


def build_content() -> list[tuple[str, str]]:
    return [
        ("h1", "进展简述"),
        ("p", "开题以来，本研究围绕家庭语音设备中“一个设备服务多个家庭成员、多个生活场景”的问题，重新梳理了研究对象、研究框架与原型方向。中期答辩后，结合评委意见，本研究将题目调整为“面向需求多样性的AI语音智能体个性化匹配与交互设计研究”。这里的“AI语音智能体”主要指用户可感知的服务角色和交互角色，而不是计算机领域中强调算法协作的 multi-agent 概念。本文关注的核心，是当家庭设备接入多个 AI 语音角色后，用户如何获得合适服务，又如何理解当前是谁在服务、为什么切换以及系统处于什么状态。"),
        ("p", "目前已完成的阶段性工作包括：第一，完成居家语音助手使用场景与用户偏好研究，归纳出休闲陪伴、健康咨询、智能家居、信息查询和学习辅导五类场景，并通过问卷量化不同用户类型在不同场景下对 AI 角色特征的偏好差异；第二，完成多智能体语音硬件的人机交互研究，明确屏幕显示、角色可见性、状态反馈和信息共享权限对用户体验的影响；第三，完成多智能体语音对话原型的核心链路，包括强意图唤醒、弱意图匹配、语音画像识别、会话保持、Coze 实时语音角色接入和异步提醒；第四，完成带屏音箱和手机端控制台的原型设计，为下一阶段实验验证提供基础。"),
        ("h1", "目录"),
        ("toc", "第1章 提出问题与研究问题\n1.1 家庭语音设备中的需求变化\n1.2 现有产品与研究的不足\n1.3 研究对象与研究问题\n第2章 设计依据研究\n2.1 研究框架\n2.2 竞品与案例分析\n2.3 用户研究\n2.4 人机交互研究\n第3章 方法构建与技术原理\n3.1 智能体角色池构建\n3.2 个性化匹配机制\n3.3 多智能体交互表达方法\n3.4 技术原理\n第4章 原型设计\n4.1 带屏智能音箱原型\n4.2 手机端控制台原型\n4.3 多智能体语音系统原型\n第5章 实验验证与研究产出\n5.1 验证目标\n5.2 实验设计\n5.3 下一阶段研究计划\n5.4 阶段性研究产出"),
        ("h1", "第1章 提出问题与研究问题"),
        ("h2", "1.1 家庭语音设备中的需求变化"),
        ("p", "智能音箱、家庭中控和手机语音助手正在进入日常生活，但家庭语音设备面对的并不是一个单一用户、一个单一任务。真实家庭中，同一台设备可能早上给老人播报天气和用药提醒，中午帮助家长查询信息，晚上陪儿童讲故事或进行英语口语练习。也就是说，家庭语音设备的使用对象、使用场景和服务期待天然具有多样性。"),
        ("p", "这种多样性带来一个很直观的生活问题：家庭成员说出的需求往往很自然、很模糊，但背后需要的服务角色并不相同。例如“我最近睡不好”可能需要健康建议，也可能需要情绪陪伴；“陪一下小孩”可能是讲故事、学习辅导，也可能是监督电视时间。对于用户来说，他们并不想先判断“应该叫哪个助手”，也不想记住每个助手的名称和能力。他们期待的是像家庭成员之间自然交流一样，系统能够理解当前需求，并给出合适的服务回应。"),
        ("p", "因此，本研究的必要性首先来自日常家庭使用中的真实痛点：单一语音助手很难同时满足不同年龄、不同任务和不同情绪语境下的期待；而多个 AI 角色虽然能够提供更细分的服务，却又可能带来新的选择负担和理解负担。设计研究需要回答的不是“能不能接入更多 AI”，而是“接入多个 AI 角色后，用户是否更容易获得合适服务，是否更清楚当前是谁在服务”。"),
        ("h2", "1.2 现有产品与研究的不足"),
        ("p", "从现有产品看，小米音箱等主流家庭语音设备通常以单一助手作为统一入口，优势是简单直接，但缺点是角色风格和服务边界较固定。当同一助手既要处理儿童陪伴，又要回答健康咨询，还要承担学习辅导和家居控制时，用户很难形成稳定的角色认知，也容易在专业性、亲和力和可信度之间产生期待落差。"),
        ("p", "亚马逊、Google、三星等产品与生态中已经出现多个语音助手共存的探索，说明产业界也意识到单一助手难以覆盖全部服务生态。然而，多助手共存并不等于体验自然。若用户仍需要记住不同唤醒词、理解不同助手能力、手动判断该找谁，系统只是把服务数量增加了，却把选择成本转移给了用户。对家庭场景而言，这种“人找服务”的方式并不适合老人、儿童和临时需求较多的日常使用。"),
        ("p", "从既有研究看，多智能体或多代理系统在计算机、移动端、网页和教育系统中已有较多讨论，但这些场景通常具有较强的视觉界面，用户可以看到多个窗口、头像或按钮。语音硬件设备则不同，用户主要通过说话完成操作，屏幕较小甚至没有屏幕。因此，家庭语音硬件中的多角色交互会遇到新的设计问题：用户如何知道当前是谁在听，谁在回答，系统为什么切换到另一个角色，不同角色是否共享了前文信息。"),
        ("h2", "1.3 研究对象与研究问题"),
        ("p", "基于上述背景，本研究将研究对象界定为“面向需求多样性的AI语音智能体系统”。这里的智能体强调的是家庭用户可感知、可区分、可调用的服务角色。每个 AI 语音智能体具有相对明确的服务场景、人格风格和语音形象，例如儿童陪伴、健康顾问、英语外教、娱乐陪伴和通用家居助手等。"),
        ("p", "本研究形成两个核心研究问题。RQ1：如何匹配合适智能体，即如何根据家庭成员特征、当前需求语义和使用场景，为用户匹配合适的 AI 语音角色，从而降低用户主动选择和记忆唤醒词的负担。RQ2：如何让多智能体交互可理解，即如何通过带屏音箱、角色头像、状态动效、语音反馈和手机端管理，让用户理解当前角色、系统状态、切换原因和信息边界，从而提升多智能体语音交互的可理解性与可控性。"),
        ("h1", "第2章 设计依据研究"),
        ("h2", "2.1 研究框架"),
        ("p", "本研究采用“提出问题—研究问题—设计依据研究—方法构建—原型设计—实验验证—研究产出”的整体框架。这样的结构把技术方法放在设计问题之后，避免将研究误解为单纯的 AI 系统开发。研究首先从家庭语音设备的使用痛点出发，说明 AI 语音智能体个性化匹配的必要性；随后明确“匹配合适智能体”和“让多智能体交互可理解”两个研究问题；再通过竞品与案例分析、用户研究和人机交互研究形成设计依据；最后构建匹配机制和交互表达方法，转化为带屏音箱、手机端控制台和语音系统原型，并通过实验验证其可用性、匹配有效性和交互可理解性。"),
        ("image", "研究框架图"),
        ("caption", "图1 研究框架图"),
        ("h2", "2.2 竞品与案例分析：明确设计机会"),
        ("p", "竞品与案例分析放在研究框架的前段，是为了让评委和读者先看到本研究与现有产品的差异。小米音箱等单一助手产品解决了“语音入口”的问题，但没有充分解决多用户、多场景下的角色适配问题；Alexa、Google Assistant、Bixby 等多生态产品说明多个服务体系可以共存，但用户仍需要理解不同平台和唤醒方式；现有多角色语音研究证明多个语音角色能够提升新颖性和吸引力，但也暴露出记忆名称、分辨当前角色和理解系统状态的困难。"),
        ("p", "因此，竞品与案例分析得到的设计机会是：家庭语音设备不应只是增加更多助手，而应从“用户主动找服务”转向“系统主动匹配服务”；不应只是让多个角色在后台存在，而应让用户看得见、听得懂、管得住这些角色。"),
        ("h2", "2.3 用户研究：证明需求差异真实存在"),
        ("p", "在用户研究部分，本研究通过定性访谈和定量问卷确认居家语音助手需求的多样性。定性访谈面向不同年龄、性别和家庭角色的用户，围绕家庭中可能使用语音助手的任务、期待和顾虑展开。结果显示，用户需求主要集中于休闲陪伴、健康咨询、智能家居、信息查询和学习辅导五类场景，而非单一的问答或设备控制。"),
        ("p", "进一步的定量问卷要求用户代入五类场景，对亲密陪伴型、鼓励赞赏型、成熟关怀型、专业可靠型和中性工具型五类语音助手角色进行偏好评分。结果表明，不同用户类型和不同场景会共同影响偏好。例如，儿童更容易偏好亲密陪伴和鼓励赞赏型角色；健康咨询场景更需要专业可靠和成熟关怀；学习辅导则需要鼓励性与专业解释并存。这一结果说明，家庭语音助手的匹配不能只看用户说了什么，还要考虑谁在说、处于什么场景，以及用户期待什么样的回应风格。"),
        ("h2", "2.4 人机交互研究：明确角色、状态与边界问题"),
        ("p", "人机交互研究部分关注的是：当多个 AI 语音智能体进入同一个家庭硬件后，用户如何理解这个系统。与手机或电脑不同，智能音箱的屏幕较小，交互以语音为主。如果多个角色都从同一个设备中说话，用户很容易产生疑问：当前是谁在回答，系统有没有听到我的话，为什么换了一个角色，前面的信息是否被其他角色知道。"),
        ("p", "基于多角色语音硬件原型的前期研究，本研究将关键交互变量归纳为外部视觉显示和信息传递边界。屏幕和头像能够增强角色存在感，帮助用户快速判断当前服务者；状态动效能够表达 listening、thinking、speaking 等过程，减少用户等待时的不确定；信息共享或旁听提示则关系到用户对隐私边界和系统可信度的理解。由此，本研究把多角色语音交互的设计目标概括为三个方面：角色易分辨、状态可感知、边界可理解。"),
        ("h1", "第3章 方法构建与技术原理"),
        ("h2", "3.1 智能体角色池构建"),
        ("p", "根据用户研究得到的五类场景，本研究构建了一个可扩展的 AI 语音智能体角色池。当前原型中包括儿童陪伴角色、健康顾问角色、英语外教角色、娱乐陪伴角色和通用家居助手角色。每个角色不仅对应不同能力，也对应不同语音风格、头像形象、适用场景和唤醒词。"),
        ("p", "角色池的设计目的不是简单增加角色数量，而是让家庭中的不同需求有更合适的服务承接者。儿童故事和健康咨询不应由同一套语气和身份完成；学习辅导和娱乐陪伴也不应采用完全相同的回应方式。通过角色池，系统能够把家庭需求拆分为更具体、更容易被用户理解的服务角色。"),
        ("h2", "3.2 个性化匹配机制"),
        ("p", "在匹配方法上，本研究将用户请求分为强意图和弱意图。强意图指用户明确说出角色名称或唤醒词，例如“叫虎子讲故事”“让健康顾问回答”；这类请求直接进入对应角色。弱意图指用户只自然表达需求，例如“我今天肚子痛”“陪小孩一会儿”“我想练英语”，系统需要根据语义和用户特征判断合适角色。"),
        ("p", "弱意图匹配采用“偏好数据召回 + 语义判断”的方式。系统首先根据用户类型、场景类型和前期问卷得到的偏好数据，对当前角色池进行初步排序；再结合用户具体表达、当前上下文和角色描述，判断最适合的 AI 角色。这样的设计既避免完全依赖关键词，也避免把选择压力全部交给模型。对于设计研究而言，它的意义在于把用户研究得到的偏好数据真正转化为系统行为。"),
        ("p", "目前阶段实现的是“场景化个性化匹配”，即系统主要根据年龄段、性别、当前语义和场景完成一次对话中的匹配。后续可以进一步拓展为长期个性化，例如记住某个家庭成员常用角色、对语气的偏好、固定任务安排和隐私边界。但在中期阶段，本研究更强调先解决“用户不用记住和选择角色”的基本体验问题。"),
        ("h2", "3.3 多智能体交互表达方法"),
        ("p", "多角色语音系统如果只在后台完成匹配，用户仍然会觉得系统像黑箱。因此，本研究把交互表达作为原型设计的重要部分。音箱屏幕用于显示当前角色头像、名称和工作状态；当系统切换角色时，界面需要给出简短的切换提示，让用户知道为什么由这个角色回答；在语音输出中，也应通过语气和称呼强化角色身份。"),
        ("p", "手机端则承担更复杂的可视化管理功能，包括查看完整角色池、了解每个角色能力、设置唤醒词、管理提醒任务和调整隐私权限。音箱端适合自然语音和轻量反馈，手机端适合复杂设置和长期管理。这样的分工可以避免把所有信息塞进小屏音箱，也能让用户在需要时获得充分控制。"),
        ("h2", "3.4 技术原理"),
        ("p", "本节集中说明原型背后的技术实现逻辑。由于本研究面向设计学问题，技术原理主要作为原型可运行和实验可验证的支撑，不作为全文的叙述主线。读者即使不深入理解具体模型，也可以通过前后章节理解本研究的设计问题、设计方法和实验验证。"),
        ("p", "系统技术链路包括语音输入、语音转写、用户画像识别、意图判断、角色匹配、AI 语音角色调用和界面状态反馈。用户说话后，系统先将语音转为文字，并在需要时通过语音片段识别年龄段和性别等粗粒度画像；随后判断用户是否明确唤醒某个角色。若用户明确点名，则直接进入对应角色；若用户表达的是模糊需求，则结合用户类型、场景判断和先验偏好数据生成候选角色，再由语义模型在候选中选择最合适的角色。"),
        ("p", "为保证实时语音体验，系统还做了耗时优化：强意图和短时间会话保持场景下不重复进行画像识别；画像模型采用本地部署的轻量语音特征模型，并通过常驻进程减少冷启动；路由模型只在弱意图场景下调用，且只在少量候选角色中判断。这样的技术处理目标不是追求复杂模型本身，而是让“个性化匹配”能够真正嵌入实时对话体验中。"),
        ("p", "语音角色调用部分主要使用 Coze 的实时语音能力，使不同角色具有不同声音、人设和回复风格。前端同步显示当前角色头像和状态，后端管理角色配置、唤醒词、偏好标签、任务记忆和隐私边界。整体技术架构服务于一个设计目标：让用户自然说出需求后，系统能够选择合适角色，并让这个选择在界面上可感知。"),
        ("h1", "第4章 原型设计"),
        ("h2", "4.1 带屏智能音箱原型"),
        ("p", "硬件原型定位为带屏智能音箱。外观方案从“显示能力”和“角色具象程度”两个维度展开，最终选择具象化外观与小屏显示结合的方向。该方向既能表达多角色陪伴的产品特征，又不会因为高度拟人化或大屏幕设计破坏家居设备的日常感。原型采用 2.8 寸圆形 IPS 屏幕呈现角色头像与状态，语音输入输出使用 Y28HM 型麦克风与扬声器模块，外壳通过 3D 打印完成。"),
        ("h2", "4.2 手机端控制台原型"),
        ("p", "手机端控制台用于承担音箱小屏难以承载的复杂管理任务，包括查看智能体角色池、了解每个角色的能力与唤醒词、管理异步提醒任务、设置隐私权限和调整家庭成员偏好。这样的分工使音箱端保持轻量自然，手机端则承担更完整的配置与解释功能。"),
        ("h2", "4.3 多智能体语音系统原型"),
        ("p", "系统原型已接入多个 Coze 语音角色，实现语音输入、文本转写、角色匹配、实时语音输出、会话保持和异步提醒。异步提醒场景用于模拟家庭中的跨时间任务，例如家长要求系统监督儿童看电视时间，到点后由更适合儿童沟通的角色进行提醒。这一功能体现了多智能体语音系统不只是即时问答工具，也可以成为家庭任务管理的一部分。"),
        ("h1", "第5章 实验验证与研究产出"),
        ("h2", "5.1 验证目标"),
        ("p", "中期答辩后，本研究对实验验证目标进行了调整。原先实验更偏向验证系统可用性，例如带屏音箱是否更好用、是否更有吸引力。但评委意见指出，仅验证可用性不足以说明个性化匹配的有效性。因此，后续实验将分为两层：第一层验证本研究原型相较普通小米音箱是否提升使用体验；第二层专门验证系统个性化匹配是否比用户手动切换更合适。"),
        ("h2", "5.2 实验设计"),
        ("p", "可用性验证部分设置两个组别：普通小米音箱与本研究带屏多角色 AI 语音助手原型。参与者将在休闲陪伴、健康咨询、智能家居、信息查询和学习辅导五类场景中完成任务。评价指标包括 SUS 可用性、认知负荷、吸引力、购买意愿和访谈反馈，用于判断本研究原型是否相较普通单助手音箱改善整体体验。"),
        ("p", "匹配有效性验证部分专门比较两种方式：用户手动切换角色与系统个性化匹配角色。在手动切换条件下，用户需要根据角色名称和能力说明自行选择；在个性化匹配条件下，用户只自然表达需求，由系统选择合适角色。每个任务完成后，参与者评价系统选择的角色是否适合当前需求、语气是否符合期待、选择过程是否省力、是否愿意继续让该角色处理类似任务。这样可以直接回应“个性化匹配是否有效”的问题，而不仅仅说明原型能运行。"),
        ("h2", "5.3 下一阶段研究计划"),
        ("p", "下一阶段工作包括四项。第一，继续完善原型系统稳定性，包括 Coze 实时语音连接、弱意图匹配准确率、任务提醒和前端状态反馈。第二，补充竞品对比图和研究框架图，使研究必要性、创新点和现有产品差异更加清晰。第三，整理实验任务、量表和访谈提纲，进行小规模预实验，检查任务描述、语音交互、角色切换和匹配评价是否清楚。第四，正式开展用户实验，并结合量化结果与访谈材料，总结面向家庭语音设备的多角色 AI 语音助手设计原则。"),
        ("h2", "5.4 阶段性研究产出"),
        ("p", "目前阶段形成的研究产出主要包括三类。第一，面向需求多样性的 AI 语音智能体个性化匹配流程，将用户特征、场景判断、偏好数据和语义判断结合起来，回应“如何匹配合适智能体”的问题。第二，带屏智能音箱、手机端控制台和多智能体语音系统原型，用于呈现和验证多角色语音交互方式。第三，初步形成角色可分辨、状态可感知、边界可理解的设计原则，为后续设计优化和论文总结提供基础。"),
        ("h1", "阶段性小结"),
        ("p", "总体来看，本研究已从最初的技术原型探索，逐步调整为更明确的设计研究：以家庭需求多样性为问题起点，以 AI 语音智能体为设计对象，以个性化匹配和可感知交互为核心方法，以带屏音箱和手机端控制台为原型载体。后续研究将重点证明两件事：一是多智能体语音系统是否比单一助手更适合家庭多场景服务；二是系统自动匹配是否真正降低用户选择负担，并提高角色匹配的合适度。"),
    ]


def main() -> None:
    make_framework_image()
    shutil.copyfile(SRC, OUT)
    doc = Document(str(OUT))

    # Update title in the school form.
    for table in doc.tables:
        for row in table.rows:
            if row.cells and "论文题目" in row.cells[0].text:
                row.cells[1].text = "面向需求多样性的AI语音智能体个性化匹配与交互设计研究"

    # Remove old report body between the report prompt and achievement section.
    start = 18
    end = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith("成果清单"):
            end = i
            break
    if end is None:
        raise RuntimeError("未找到成果清单段落")
    for p in doc.paragraphs[start:end]:
        p._element.getparent().remove(p._element)

    achievement = None
    for p in doc.paragraphs:
        if p.text.startswith("成果清单"):
            achievement = p
            break
    if achievement is None:
        raise RuntimeError("未找到成果清单段落")

    for typ, text in build_content():
        if typ == "h1":
            add_para(achievement, text, "Report Heading 1")
        elif typ == "h2":
            add_para(achievement, text, "Report Heading 2")
        elif typ == "h3":
            add_para(achievement, text, "Report Heading 3")
        elif typ == "caption":
            add_para(achievement, text, "Report Caption", WD_ALIGN_PARAGRAPH.CENTER)
        elif typ == "toc":
            for line in text.splitlines():
                add_para(achievement, line, "Report TOC")
        elif typ == "image":
            add_image(achievement, IMG, 6.6)
        else:
            add_para(achievement, text, "Report Body")

    doc.save(str(OUT))
    print(OUT)
    print(IMG)


if __name__ == "__main__":
    main()
