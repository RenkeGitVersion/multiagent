from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOCX = Path("/Users/renke/毕设/中期答辩/任科中期报告_0701_新框架版.docx")


def insert_after(paragraph: Paragraph, text: str, style: str = "Report Caption") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = style
    inserted.add_run(text)
    return inserted


def main() -> None:
    doc = Document(str(DOCX))
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "图 新体验二：不同用户的差异化匹配":
            paragraph.text = "图 新体验一：基于需求的智能唤醒"
            break

    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().startswith("第二，同一输入在不同用户"):
            for candidate in doc.paragraphs[i + 1 : i + 5]:
                if candidate.text.strip() == "图 新体验二：不同用户的差异化匹配":
                    return
                if "w:drawing" in candidate._p.xml:
                    insert_after(candidate, "图 新体验二：不同用户的差异化匹配")
                    break
            break

    doc.save(str(DOCX))
    print(DOCX)


if __name__ == "__main__":
    main()
