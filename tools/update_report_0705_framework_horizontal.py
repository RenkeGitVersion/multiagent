from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


DOCX = Path("/Users/renke/毕设/中期答辩/任科中期报告_0701_新框架版.docx")
OUT_IMG = Path("/Users/renke/Documents/Codex/2026-05-27/1-1-2-agent-1-agent/tmp/research_framework_0705_horizontal.png")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = "/System/Library/Fonts/PingFang.ttc"
    if Path(path).exists():
        return ImageFont.truetype(path, size=size, index=1 if bold else 0)
    return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines = []
    for raw in text.split("\n"):
        line = ""
        for ch in raw:
            test = line + ch
            if draw.textbbox((0, 0), test, font=fnt)[2] <= max_width:
                line = test
            else:
                if line:
                    lines.append(line)
                line = ch
        if line:
            lines.append(line)
    return lines


def box(draw, xy, title, body, fill, outline, title_color):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=24, fill=fill, outline=outline, width=4)
    tf, bf = font(31, True), font(22)
    title_lines = wrap(draw, title, tf, x2 - x1 - 44)
    body_lines = wrap(draw, body, bf, x2 - x1 - 48)
    total_h = len(title_lines) * 39 + 8 + len(body_lines) * 29
    y = y1 + (y2 - y1 - total_h) / 2
    for line in title_lines:
        w = draw.textbbox((0, 0), line, font=tf)[2]
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=tf, fill=title_color)
        y += 39
    y += 8
    for line in body_lines:
        w = draw.textbbox((0, 0), line, font=bf)[2]
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=bf, fill="#374151")
        y += 29


def arrow(draw, start, end, color="#475569"):
    draw.line([start, end], fill=color, width=5)
    sx, sy = start
    ex, ey = end
    if ex > sx:
        pts = [(ex, ey), (ex - 24, ey - 13), (ex - 24, ey + 13)]
    elif ey > sy:
        pts = [(ex, ey), (ex - 13, ey - 24), (ex + 13, ey - 24)]
    else:
        pts = [(ex, ey), (ex - 13, ey + 24), (ex + 13, ey + 24)]
    draw.polygon(pts, fill=color)


def main() -> None:
    OUT_IMG.parent.mkdir(parents=True, exist_ok=True)
    W, H = 2400, 1450
    img = Image.new("RGB", (W, H), "#f8fafc")
    draw = ImageDraw.Draw(img)
    draw.text((W / 2, 70), "研究框架图", anchor="mm", font=font(58, True), fill="#111827")
    draw.text((W / 2, 122), "面向需求多样性的 AI 语音智能体个性化匹配与交互设计研究", anchor="mm", font=font(26), fill="#4b5563")

    # Column headings
    for x, label in [(360, "问题提出"), (1050, "研究与方法构建"), (1860, "原型与验证")]:
        draw.text((x, 190), label, anchor="mm", font=font(32, True), fill="#334155")

    # Left column
    problem = (110, 245, 610, 535)
    rq = (110, 625, 610, 875)
    box(draw, problem, "第1章\n研究背景与问题", "家庭多用户、多场景需求多样\n单一智能体适配不足\n多智能体带来服务发现与理解问题", "#fff7ed", "#ea580c", "#7c2d12")
    box(draw, rq, "研究问题", "RQ1 如何动态匹配合适智能体？\nRQ2 如何让多智能体交互可理解？", "#eef2ff", "#4f46e5", "#312e81")
    arrow(draw, (360, 535), (360, 625))

    # Middle column top
    ch2 = (760, 245, 1340, 455)
    ch3 = (760, 525, 1340, 735)
    box(draw, ch2, "第2章\n相关研究与竞品分析", "文献研究：需求多样性、单智能体局限、多智能体优势\n竞品调研：传统音箱与新型多智能体音箱", "#ecfeff", "#0891b2", "#164e63")
    box(draw, ch3, "第3章\n体验设计框架构建", "用户研究 → 用户体验重构\n个性化匹配设计框架", "#f0fdf4", "#16a34a", "#14532d")
    arrow(draw, (610, 750), (760, 635))
    arrow(draw, (1050, 455), (1050, 525))

    # Middle split: RQ1/RQ2
    ch4 = (680, 835, 1240, 1120)
    ch5 = (1320, 835, 1880, 1120)
    box(draw, ch4, "第4章\n个性化匹配方法构建", "对应 RQ1\n即时情境响应 → 用户偏好 → 长期记忆\n三层递进形成完整匹配链路", "#eff6ff", "#2563eb", "#1e3a8a")
    box(draw, ch5, "第5章\n多智能体交互设计研究", "对应 RQ2\n外观与交互形象研究\n屏幕显示 × 信息共享实验\n形成可理解交互原则", "#faf5ff", "#9333ea", "#581c87")
    arrow(draw, (1050, 735), (960, 835))
    arrow(draw, (1050, 735), (1510, 835))

    # Right column
    ch6 = (1960, 275, 2290, 505)
    ch7 = (1960, 610, 2290, 840)
    ch8 = (1960, 945, 2290, 1175)
    box(draw, ch6, "第6章\n原型设计与开发", "系统设计与开发\n音箱设计与开发\n手机端设计与开发", "#fff7ed", "#ea580c", "#7c2d12")
    box(draw, ch7, "第7章\n实验验证", "原型用户体验验证\n方法有效性验证\n个性化匹配 vs 唤醒词/手动切换", "#fef2f2", "#dc2626", "#7f1d1d")
    box(draw, ch8, "第8章\n研究总结", "理论贡献\n方法贡献\n设计贡献\n实践贡献", "#f8fafc", "#475569", "#111827")
    arrow(draw, (1880, 980), (1960, 390))
    arrow(draw, (1960, 390), (1880, 980), "#475569")
    # clean overloaded diagonal by drawing central merge
    draw.line([(1240, 980), (1580, 980), (1580, 390), (1960, 390)], fill="#475569", width=5)
    draw.line([(1880, 980), (1960, 390)], fill="#f8fafc", width=10)
    draw.polygon([(1960, 390), (1936, 377), (1936, 403)], fill="#475569")
    arrow(draw, (2125, 505), (2125, 610))
    arrow(draw, (2125, 840), (2125, 945))

    # bottom logic rail
    rail_y = 1310
    draw.rounded_rectangle((180, rail_y - 44, W - 180, rail_y + 44), radius=22, fill="#ffffff", outline="#cbd5e1", width=3)
    draw.text(
        (W / 2, rail_y),
        "逻辑主线：发现必要性 → 建立体验框架 → 构建匹配方法与可理解交互 → 开发原型 → 实验验证 → 总结贡献",
        anchor="mm",
        font=font(27, True),
        fill="#334155",
    )
    img.save(OUT_IMG, quality=95)

    doc = Document(str(DOCX))
    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "图1 研究框架图" and i > 0:
            img_para = doc.paragraphs[i - 1]
            img_para.clear()
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_para.add_run().add_picture(str(OUT_IMG), width=Inches(6.4))
            break
    doc.save(str(DOCX))
    print(DOCX)
    print(OUT_IMG)


if __name__ == "__main__":
    main()
