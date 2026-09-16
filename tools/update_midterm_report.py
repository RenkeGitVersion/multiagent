from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, Cm, Inches


SRC = Path("/Users/renke/毕设/中期答辩/任科中期报告.docx")
BACKUP = SRC.with_name("任科中期报告_备份_修改前.docx")
SLIDE_DIR = Path("/Users/renke/Documents/Codex/2026-05-27/1-1-2-agent-1-agent/tmp/midterm_slides")


def set_run_font(run, east="仿宋_GB2312", ascii_font="Times New Roman", size=12, bold=False):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold


def ensure_styles(doc):
    styles = doc.styles
    for name, east, size, bold in [
        ("Report Body", "仿宋_GB2312", 12, False),
        ("Report Caption", "仿宋_GB2312", 10.5, False),
    ]:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH) if name not in styles else styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), east)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.25

    styles["Report Body"].paragraph_format.first_line_indent = Cm(0.74)
    styles["Report Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Report Caption"].paragraph_format.space_after = Pt(8)

    for name, east, size in [
        ("Report Heading 1", "黑体", 15),
        ("Report Heading 2", "黑体", 13),
        ("Report Heading 3", "黑体", 12),
    ]:
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH) if name not in styles else styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), east)
        st.font.size = Pt(size)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(8 if name != "Report Heading 1" else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.25


def main():
    if not BACKUP.exists():
        shutil.copy2(SRC, BACKUP)

    # Rebuild from the original template each time so reruns do not duplicate the report body.
    doc = Document(BACKUP)
    ensure_styles(doc)
    styles = doc.styles

    start_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith("报告正文 Report"))
    end_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith("成果清单 List of Achievements"))
    start_el = doc.paragraphs[start_idx]._p
    end_el = doc.paragraphs[end_idx]._p
    body = start_el.getparent()

    for p in list(doc.paragraphs[start_idx + 1:end_idx]):
        body.remove(p._p)

    def paragraph_from_element(element):
        for candidate in doc.paragraphs:
            if candidate._p is element:
                return candidate
        raise RuntimeError("paragraph not found")

    def insert_paragraph(text="", style="Report Body"):
        p = OxmlElement("w:p")
        body.insert(body.index(end_el), p)
        para = paragraph_from_element(p)
        if style:
            para.style = styles[style] if style in styles else style
        if text:
            run = para.add_run(text)
            if style == "Heading 1":
                set_run_font(run, "黑体", size=15, bold=True)
            elif style == "Heading 2":
                set_run_font(run, "黑体", size=13, bold=True)
            elif style == "Heading 3":
                set_run_font(run, "黑体", size=12, bold=True)
            elif style == "Report Caption":
                set_run_font(run, "仿宋_GB2312", size=10.5)
            else:
                set_run_font(run, "仿宋_GB2312", size=12)
        return para

    def heading(text, level=1):
        p = insert_paragraph("", f"Report Heading {level}")
        r = p.add_run(text)
        set_run_font(r, "黑体", size=15 if level == 1 else 13 if level == 2 else 12, bold=True)
        return p

    def body_text(text):
        p = insert_paragraph("", "Report Body")
        r = p.add_run(text)
        set_run_font(r, "仿宋_GB2312", size=12)
        return p

    def caption(text):
        p = insert_paragraph("", "Report Caption")
        r = p.add_run(text)
        set_run_font(r, "仿宋_GB2312", size=10.5)
        return p

    def picture(slide_no, text, width=5.35):
        img = SLIDE_DIR / f"slide-{slide_no:02d}.png"
        if not img.exists():
            return
        p = insert_paragraph("", None)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img), width=Inches(width))
        caption(text)

    heading("目录", 1)
    for line in [
        "一、研究背景与问题提出",
        "二、开题以来的研究进展",
        "    （一）用户研究：居家语音智能体需求与偏好",
        "    （二）人机交互研究：多智能体语音硬件的交互原则",
        "    （三）技术方法研究：个性化匹配与多智能体语音对话架构",
        "三、阶段性成果",
        "四、原型设计进展",
        "五、后续实验设计与下一阶段计划",
    ]:
        p = insert_paragraph("", "Report Body")
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(line)
        set_run_font(r, "仿宋_GB2312", size=12)

    heading("一、研究背景与问题提出", 1)
    for text in [
        "本课题围绕“面向需求多样性的 AI 语音智能体个性化匹配与交互设计”展开。开题以来，研究重点从“单一语音助手能否满足家庭场景需求”转向“多智能体如何在同一家庭公共设备中被理解、被选择和被管理”。随着大语言模型和语音交互能力的发展，智能体正在从个人手机、电脑等私有设备，进入智能音箱、家庭中控等公共设备。家庭空间中的用户具有不同年龄、身份和认知习惯，同一设备也需要同时覆盖学习、陪伴、健康咨询、信息查询和家居控制等任务。因此，居家语音智能体面对的并不是单一用户、单一任务，而是多用户、多场景、多偏好的混合情境。",
        "现有单智能体语音助手通常采用相对固定的声音、角色和服务方式，在简单问答或设备控制中可以工作，但在复杂家庭场景中容易出现角色冲突。例如，同一个助手既要讲故事、陪儿童聊天，又要提供健康咨询和学习辅导时，用户对其专业性、亲和力和可信度的期待并不相同。相比之下，多智能体系统能够通过角色划分，让不同智能体承担不同类型的任务，但现有多智能体系统往往依赖用户主动记住唤醒词并手动切换，这会增加使用门槛，也容易破坏连续对话。",
        "基于上述背景，本研究形成两个主要研究问题：第一，如何构建 AI 语音智能体的个性化匹配机制，以应对不同用户、不同场景下的偏好和需求差异；第二，如何建立高效、自然、可理解的人机交互方式，使多智能体语音系统能够在硬件设备上被用户清楚地感知和使用。围绕这两个问题，开题以来的工作主要包括用户研究、人机交互研究、技术方法研究和原型设计四个部分。",
    ]:
        body_text(text)
    picture(5, "图1 研究框架：用户研究、人机交互研究、技术方法研究与原型设计之间的关系")

    heading("二、开题以来的研究进展", 1)
    heading("（一）用户研究：居家语音智能体需求与偏好", 2)
    for text in [
        "用户研究部分首先通过访谈了解不同年龄和性别用户在居家场景下对语音智能体的需求。访谈对象覆盖儿童、大学生、职场成人和中老年用户，共 12 人。主题编码结果显示，当前用户对居家语音智能体的需求主要集中在非生产力场景，尤其是休闲陪伴、健康咨询、智能家居、信息查询和学习辅导五类。办公、深度创作等任务虽然也可能使用智能体，但受限于输入效率、输出可控性和工作习惯，用户仍更倾向于使用电脑等传统设备。",
        "访谈还表明，居家语音智能体需求受到两类因素影响：一是用户特征，包括年龄、身份、认知能力和情感偏好；二是任务场景，包括陪伴、学习、健康、控制和查询等不同任务目标。不同用户在同一场景中会形成不同期待，同一用户在不同场景下也会改变对智能体角色的要求。这说明居家语音智能体的设计不能只依赖一个通用角色，而需要考虑多智能体分工和动态匹配。",
        "在定量研究中，本研究基于前期访谈总结出的五类场景，设计了用户偏好问卷。问卷要求用户分别代入休闲陪伴、健康咨询、智能家居、信息查询、学习辅导五类场景，并对亲密陪伴型、鼓励赞赏型、成熟关怀型、专业可靠型和中性工具型五类智能体给出 1-7 分偏好评分。最终共收集 131 份数据，经过直线作答、极端同分、完成时长过短和注意力检测等规则筛查后，保留 125 份有效数据。",
        "定量结果显示，不同用户类型和不同场景均会影响智能体类型偏好。总体上，儿童更偏好亲密陪伴型和鼓励赞赏型，大学生与职场成人更偏好专业可靠型，中老年用户更偏好成熟关怀型和中性工具型。从场景看，休闲陪伴更适合亲密陪伴型，健康咨询更需要专业可靠型与成熟关怀型，智能家居和信息查询更偏向专业可靠型或中性工具型，学习辅导则更重视鼓励和专业解释。性别维度下的差异相对不明显，因此在后续技术方法中，性别更多作为辅助特征，而不是主导判断依据。",
        "进一步分析“用户类型 × 场景”的交互关系后可以看到，用户类型和场景并不是相互独立地影响偏好，而是共同决定主导智能体类型。例如，儿童在学习辅导中更偏好鼓励赞赏型，而在健康咨询中则更倾向于成熟关怀和专业可靠的表达方式。这部分数据为后续多智能体角色划分、场景识别和个性化匹配提供了先验依据。",
    ]:
        body_text(text)
    picture(10, "图2 定量研究结论：用户类型与场景共同影响智能体类型偏好")

    heading("（二）人机交互研究：多智能体语音硬件的交互原则", 2)
    for text in [
        "在用户研究明确“需求多样性”之后，本研究进一步关注多智能体语音硬件的交互方式。过去的多智能体研究多集中在电脑或移动端等视觉界面上，而智能音箱等语音硬件具有不同特点：用户通常以语音为主要入口，智能体之间需要顺序发声，用户很难像在屏幕界面中那样同时看到多个智能体状态。因此，多智能体语音硬件容易出现身份识别困难、切换过程不清楚和认知负荷增加等问题。",
        "为探讨这些问题，本研究基于 AgentsNest 多智能体原型进行了受控实验，考察外部视觉显示（有屏幕与无屏幕）以及智能体间信息传输模式（共享与隔离）对用户体验的影响。实验结果显示，屏幕的存在能够提升系统易用性、减轻认知负荷、提高识别效率，并增强用户对系统的信任感。对于多智能体语音设备而言，屏幕并不一定需要承担复杂触控操作，但需要帮助用户理解当前是谁在说话、是否正在聆听、为什么发生切换，以及系统当前处于何种状态。",
        "信息共享模式的结果较为复杂。一部分用户认为智能体之间共享上下文会让系统显得更统一、更聪明，也能减少重复说明；另一部分用户则担心多个智能体同时“旁听”会带来隐私压力。由此可见，多智能体系统需要在流畅性和隐私感知之间取得平衡。一方面，适当的信息共享有助于维持对话连续性；另一方面，系统需要通过界面明确呈现哪些智能体正在参与、哪些智能体具有旁听权限，并允许用户调整这些设置。",
        "基于实验结果，研究形成了若干设计启示。第一，多智能体硬件设备最好配备视觉用户界面，尤其当智能体数量超过三个时，用户仅凭声音很难区分不同角色。第二，界面应充分反映智能体特征，例如头像、名称、性格和能力标签，帮助用户建立稳定的角色认知。第三，系统应提供信息共享或旁听状态的可视化表达，并允许用户在设备在线或离线时调整相关权限。第四，用户普遍不希望记住大量唤醒词，而是希望系统像“中心化管家”一样，根据需求自动分发任务。",
    ]:
        body_text(text)
    picture(16, "图3 人机交互研究得到的设计启示：屏幕、旁听、权限与自动匹配")

    heading("（三）技术方法研究：个性化匹配与多智能体语音对话架构", 2)
    for text in [
        "技术方法研究主要回应“如何让系统自动匹配合适智能体”这一问题。本研究参考多语音助手对话与仲裁相关白皮书中的基础语音对话状态，将语音交互过程理解为空闲、聆听、思考、说话和等待输入等状态。传统多语音助手仲裁主要基于唤醒词，即用户明确呼叫某个助手后系统进行响应。这种方式适合强意图场景，但无法处理“我今天肚子痛”“我想学英语”“帮我监督孩子看电视时间”这类没有明确唤醒词的弱意图表达。",
        "为突破单一唤醒词限制，本研究在架构中引入个性化匹配机制。系统首先区分强意图和弱意图：当用户明确说出某个智能体名称或唤醒词时，系统直接调用对应智能体；当用户没有明确指定智能体时，系统将进入弱意图匹配流程。弱意图匹配不直接让大模型在所有智能体中选择，而是先依据用户画像、场景类型和先验偏好数据召回 Top3 候选，再由 LLM 在少量候选中做最终语义判断。这样可以减少大模型盲选带来的不稳定，也使匹配过程更容易解释。",
        "具体而言，用户语音首先经过 ASR 得到 query，同时通过语音画像模块识别年龄段和性别。系统再将 query 归一到五类场景之一，包括休闲陪伴、健康咨询、智能家居、信息查询和学习辅导。随后，系统查找“用户类型 × 场景 × 智能体类型”的偏好矩阵，并为当前智能体池中的每个 agent 计算偏好分。每个 agent 预先标注两个智能体类型标签，例如健康顾问对应“专业可靠型 + 成熟关怀型”，虎子对应“亲密陪伴型 + 鼓励赞赏型”。agent 的偏好分由其两个标签在当前用户类型和场景下的平均分得到。",
        "以“儿童 + 我今天肚子痛”为例，系统首先判断用户类型为儿童，场景归一为健康咨询。根据最终问卷数据，在“儿童 × 健康咨询”下，五类智能体类型的平均偏好分分别为：亲密陪伴型 4.24、鼓励赞赏型 4.56、成熟关怀型 4.92、专业可靠型 4.48、中性工具型 3.16。由此可计算出健康顾问 4.70、孙悟空 4.58、英语外教 4.52、虎子 4.40、豆包 3.82。系统据此召回健康顾问、孙悟空和英语外教三个候选，再由 LLM 根据“肚子痛”这一具体语义选择健康顾问。这个例子说明，用户年龄会影响偏好，但具体任务场景仍然是匹配的重要依据。",
        "语音画像模块用于补充弱意图匹配中的用户特征。初版方法采用较大的 wav2vec2 年龄性别识别模型，每次请求都需要上传音频、启动 Python 进程、加载模型、推理和退出，耗时约 12-17 秒，难以用于实时交互。后续优化中，本研究将模型改为常驻 Python worker，首次加载后复用同一模型实例；同时使用精简模型、本地 Hugging Face snapshot 离线加载、音频静音裁剪、音量归一化和多段推理综合决策。优化后冷启动约 3.4 秒，热启动约 0.29 秒。为降低误判，系统还引入性别置信度阈值，低于阈值时输出 unknown，而不是强行分类。",
    ]:
        body_text(text)
    picture(19, "图4 个性化匹配链路：用户特征、场景特征、先验偏好数据与 LLM 语义匹配")
    picture(20, "图5 弱意图匹配流程：先验偏好召回排序与 LLM Top3 最终裁决")
    picture(21, "图6 性能优化：语音画像模块与路由模块的耗时优化")

    heading("三、阶段性成果", 1)
    for text in [
        "开题以来，本研究已完成三个方面的阶段性成果。第一，完成了用户研究与定量偏好分析，明确了居家语音智能体的五类主要场景、五类智能体类型以及用户类型与场景交互影响下的偏好差异。这部分结果为后续原型设计和技术方法提供了数据基础。第二，完成了多智能体语音硬件交互研究，初步验证了屏幕显示、智能体身份可见性和信息共享权限对于用户体验的重要性，并形成了针对硬件音箱和手机端控制台的设计启示。第三，完成了多智能体语音对话 demo 的核心技术链路，包括强意图唤醒、弱意图路由、会话保持、Coze 实时语音智能体接入、语音画像识别和异步任务提醒等功能。",
        "目前原型系统采用浏览器 Web Demo 与本地 Node 后端实现。前端提供录音、转写、当前 agent 显示、头像/动画展示、对话流和画像面板；后端包括 Agent Registry、Router、Coze Adapter、Profile Provider 和 Task Memory 等模块。系统已接入虎子、健康顾问、英语外教、孙悟空、豆包等五个具有不同人设和服务场景的智能体。对于强意图，系统直接按唤醒词或名称路由；对于弱意图，系统结合用户画像、场景判断和 LLM 输出进行匹配。为了改善连续对话体验，系统还加入了会话保持和自动断开逻辑，使用户在短时间内继续追问时不必重复唤醒同一智能体。",
        "在交互层面，原型验证了带屏音箱和手机端的分工：音箱侧承担自然语音入口、当前 agent 显示和主动提醒，手机端承担完整 agent 池展示、能力说明、手动切换、任务管理和隐私设置。该分工与前期用户研究和人机交互研究结论一致，即用户在实时语音对话中需要低负担入口，而在管理和设置中需要更完整的可视化界面。",
    ]:
        body_text(text)
    picture(25, "图7 系统架构图：多智能体语音交互、路由、Memory 与 Coze Realtime 连接")

    heading("四、原型设计进展", 1)
    for text in [
        "原型设计部分围绕硬件音箱、音箱交互和手机端 App 三个层面展开。硬件音箱定位为家庭空间中的自然语音入口，同时配备屏幕用于显示当前 agent、头像/动画、聆听状态、路由结果和提醒信息。与无屏音箱相比，带屏音箱可以帮助用户理解系统状态，尤其是在多个智能体轮流发声时，屏幕能够降低用户辨认“谁在说话”的负担。",
        "音箱交互设计参考了多智能体语音交互指南中对注意力状态的定义，重点呈现 listening、thinking 和 speaking 三类状态。用户说话时，音箱需要反馈正在聆听；系统处理请求时，需要显示正在思考或路由；agent 回复时，需要明确显示当前 agent 的头像、名称和声音状态。对于弱意图自动路由，界面还应在不打断对话的前提下简要说明切换原因，例如“已为你切换到健康顾问”。",
        "手机端 App 主要承担了解能力、实时交互辅助、异步任务管理和后台设置四类功能。在“了解能力”阶段，用户可以浏览完整 agent 池，查看每个智能体的人设、能力、适用场景和唤醒词；在“实时交互”阶段，手机端可显示当前 agent、路由原因、声音画像和手动切换入口；在“异步交互”阶段，用户可以查看、修改或取消提醒；在“后台设置”阶段，用户可以接入或替换智能体，设置头像/GIF、唤醒词、数据共享策略和隐私权限。",
        "基于前期研究结论，本研究将用户旅程划分为了解能力、实时交互、异步交互和后台设置四个阶段。该旅程不是只关注一次问答，而是覆盖用户从认识系统、使用系统到长期管理系统的完整过程。这样可以更好地回应多智能体系统中的记忆负担、可见性、自动匹配和隐私管理问题。",
    ]:
        body_text(text)
    picture(24, "图8 用户旅程图：了解能力、实时交互、异步交互与后台设置")

    heading("五、后续实验设计与下一阶段计划", 1)
    for text in [
        "下一阶段将围绕原型系统开展用户实验，检验多智能体个性化匹配和硬件界面设计对可用性、吸引力、购买意愿和认知负荷的影响。实验计划采用被试内设计，让参与者在五类任务场景中分别与不同系统条件进行交互，包括普通单智能体音箱、普通音箱形态下的多智能体系统，以及带屏原型音箱下的多智能体系统。核心任务覆盖休闲陪伴、健康咨询、智能家居、信息查询和学习辅导。",
        "实验评估将结合量表和访谈。量表部分拟使用 SUS 系统可用性量表、AttrakDiff 吸引力量表、购买意愿量表和 NASA-TLX 认知负荷量表。访谈部分将重点询问用户是否能感知不同智能体分工、是否存在记忆唤醒词的负担、是否信任自动路由、是否需要屏幕辅助理解系统状态，以及对信息共享或旁听机制的隐私感受。实验假设包括：相比普通单智能体音箱，多智能体个性化匹配音箱能够提升系统可用性、吸引力和购买意愿；相比普通音箱的外观和显示方式，带屏原型音箱能够进一步降低认知负荷，并提高用户对多智能体状态的理解。",
        "后续工作将分为三步推进。第一，继续完善原型系统的稳定性，包括 Coze 实时语音连接、弱意图路由准确率、任务提醒和前端状态反馈。第二，完成实验材料、任务脚本和量表问卷的整理，并进行小规模预测试，检查任务难度和系统稳定性。第三，正式开展用户实验，对量表结果进行统计分析，并结合访谈材料总结多智能体语音硬件的设计原则。最终论文将围绕“需求多样性”“个性化匹配机制”“硬件交互可理解性”和“用户可控的隐私管理”四个方面展开。",
    ]:
        body_text(text)
    picture(29, "图9 后续实验设计：实验组别、任务场景、评价量表与研究假设")

    p = insert_paragraph("", None)
    p.add_run().add_break()

    doc.save(SRC)
    print(SRC)
    print(BACKUP)


if __name__ == "__main__":
    main()
