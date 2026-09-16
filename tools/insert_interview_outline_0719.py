from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


DOCX = Path("/Users/renke/毕设/中期答辩/任科中期报告_0719_4-3长期记忆补文版.docx")


ROWS = [
    (
        "智能体感知与自动路由",
        "分工感知",
        "体验多智能体时，你能明显感到是不同“专家”在服务，还是觉得它们其实是一回事？",
    ),
    (
        "智能体感知与自动路由",
        "记忆负担",
        "场景切换时，需要特意去记哪个 Agent 对应什么功能吗？有没有“记不住名字、不知道该喊谁”的时候？",
    ),
    (
        "智能体感知与自动路由",
        "自动路由",
        "如果系统能根据你的话自动转接最合适的智能体，你觉得体验有提升吗？你更信任自动分配，还是倾向手动切换？",
    ),
    (
        "视觉界面的必要性",
        "纯声音困境",
        "在无屏小米音箱上，仅靠换声线，你能分清是谁在说话吗？会不会有“看不见后台”的感觉？",
    ),
    (
        "视觉界面的必要性",
        "视觉释放负荷",
        "换到带屏原型后，屏幕上的头像和动态有减轻你“猜测谁在说话”的脑力负担吗？",
    ),
    (
        "视觉界面的必要性",
        "触控与显示",
        "你觉得这类音箱必须配像手机一样的触控屏，还是只要一个看状态、露脸的普通显示屏就够了？",
    ),
    (
        "UI 特征与旁听心理",
        "视觉拟人化",
        "视觉 UI 和头像符合你听到声音时对它们“性格、专业度”的想象吗？",
    ),
    (
        "UI 特征与旁听心理",
        "旁听心理安全感",
        "当 A 在说话而 B 在“旁听”协同时，界面用眼睛睁闭提示，你会觉得有掌控感，还是会担心隐私、觉得被监视？",
    ),
    (
        "综合总结",
        "对比与购买",
        "对比普通音箱，你更信任哪一个？如果新原型更贵，你愿意为这个“多智能体带屏”体验买单吗？最吸引你的点是什么？",
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_row_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def format_cell(cell, bold=False, color=None, size=8.5) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(size)
            run.bold = bold
            if color:
                run.font.color.rgb = color


def insert_paragraph_before(paragraph: Paragraph, text: str, style: str = "Report Body") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = style
    inserted.add_run(text)
    return inserted


def main() -> None:
    doc = Document(str(DOCX))

    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip() in {
            "表 原型用户体验验证访谈大纲",
            "在完成任务操作和量表评分后，研究将进一步进行半结构化访谈，用于补充解释用户在多智能体感知、自动路由、视觉界面和隐私心理等方面的主观体验。",
        }:
            paragraph._element.getparent().remove(paragraph._element)

    for table in list(doc.tables):
        if table.rows and table.rows[0].cells[0].text.strip() == "访谈维度":
            table._tbl.getparent().remove(table._tbl)

    target = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("7.2 方法有效性验证") and paragraph.style.name == "Report Heading 2":
            target = paragraph
            break
    if target is None:
        raise ValueError("7.2 方法有效性验证 not found")

    intro = insert_paragraph_before(
        target,
        "在完成任务操作和量表评分后，研究将进一步进行半结构化访谈，用于补充解释用户在多智能体感知、自动路由、视觉界面和隐私心理等方面的主观体验。",
    )
    caption = insert_paragraph_before(target, "表 原型用户体验验证访谈大纲", "Report Caption")

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table._tbl.getparent().remove(table._tbl)
    target._p.addprevious(table._tbl)

    widths = [2100, 1900, 5400]
    headers = ["访谈维度", "关注点", "访谈问题"]
    set_row_repeat_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, "1F5DFF")
        set_cell_margins(cell)
        format_cell(cell, bold=True, size=8.5)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)

    for dimension, focus, question in ROWS:
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        values = [dimension, focus, question]
        for i, value in enumerate(values):
            cells[i].text = value
            set_cell_width(cells[i], widths[i])
            set_cell_margins(cells[i])
            format_cell(cells[i], size=8.5)
        if dimension in {"智能体感知与自动路由", "UI 特征与旁听心理"}:
            for cell in cells:
                set_cell_shading(cell, "F7FAFF")

    doc.save(str(DOCX))
    print(DOCX)


if __name__ == "__main__":
    main()
