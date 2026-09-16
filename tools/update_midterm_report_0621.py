from __future__ import annotations

from pathlib import Path
import re
import shutil
import subprocess

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


PROJECT_ROOT = Path("/Users/renke/Documents/Codex/2026-05-27/1-1-2-agent-1-agent")
PYTHON = Path("/Users/renke/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3")
RENDER = Path("/Users/renke/.codex/plugins/cache/openai-primary-runtime/documents/26.619.11828/skills/documents/render_docx.py")
PDFTOTEXT = Path("/Users/renke/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/poppler/bin/pdftotext")
SRC = Path("/Users/renke/毕设/中期答辩/任科中期报告_0621.docx")
BACKUP = Path("/Users/renke/毕设/中期答辩/任科中期报告_0621_备份_修改前.docx")
TEMP_DIR = PROJECT_ROOT / "tmp" / "midterm_report_0621_toc"


TOC_ITEMS = [
    (1, "第1章 研究背景与研究问题", "研究背景与研究问题"),
    (2, "研究背景", "1.1 研究背景"),
    (2, "研究现状", "1.2 研究现状"),
    (2, "现有研究的局限性", "1.3 现有研究的局限性"),
    (2, "研究问题", "1.4 研究问题"),
    (1, "第2章 研究内容", "研究内容"),
    (2, "2.1 研究框架", "2.1 研究框架"),
    (2, "2.2 用户研究", "2.2 用户研究"),
    (2, "2.3 人机交互研究", "2.3 人机交互研究"),
    (2, "2.4 技术方法研究", "2.4 技术方法研究"),
    (1, "第3章 原型设计", "原型设计"),
    (2, "3.1 设计策略", "3.1 设计策略"),
    (2, "3.2 用户旅程", "3.2 用户旅程"),
    (2, "3.3 系统架构", "3.3 系统架构"),
    (2, "3.4 硬件设计", "3.4 硬件设计"),
    (2, "3.5 手机端APP设计", "3.5 手机端APP设计"),
    (1, "第4章 实验验证", "实验验证"),
    (2, "4.1 实验设计", "4.1 实验设计"),
    (1, "第5章 下一阶段研究计划", "下一阶段研究计划"),
]


def norm(text: str) -> str:
    return " ".join(text.split())


def ensure_style(doc: Document, name: str, east: str, size: float, bold: bool = False, first_line: bool = False):
    styles = doc.styles
    style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH) if name not in styles else styles[name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.first_line_indent = Cm(0.74) if first_line else Cm(0)
    return style


def set_run_font(run, east="仿宋_GB2312", ascii_font="Times New Roman", size=12, bold=False):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold


def insert_after(paragraph, text: str, style_name: str):
    parent = paragraph._p.getparent()
    idx = parent.index(paragraph._p) + 1
    from docx.oxml import OxmlElement

    p_el = OxmlElement("w:p")
    parent.insert(idx, p_el)
    doc = paragraph.part.document
    para = next(p for p in doc.paragraphs if p._p is p_el)
    para.style = doc.styles[style_name]
    run = para.add_run(text)
    set_run_font(run)
    return para


def delete_paragraph(paragraph):
    p = paragraph._p
    p.getparent().remove(p)


def apply_heading_styles(doc: Document):
    for p in doc.paragraphs:
        t = norm(p.text)
        if not t:
            continue
        if t in {"研究背景与研究问题", "第1章 研究背景与研究问题", "研究内容", "第2章 研究内容", "原型设计", "第3章 原型设计", "实验验证", "第4章 实验验证", "下一阶段研究计划", "第5章 下一阶段研究计划"}:
            p.style = doc.styles["Report Heading 1"]
            p.paragraph_format.first_line_indent = Cm(0)
        elif re.fullmatch(r"1\.\d .+", t) or re.fullmatch(r"[234]\.\d .+", t):
            p.style = doc.styles["Report Heading 2"]
            p.paragraph_format.first_line_indent = Cm(0)
        elif re.fullmatch(r"[234]\.\d\.\d .+", t):
            p.style = doc.styles["Report Heading 3"]
            p.paragraph_format.first_line_indent = Cm(0)


def replace_toc(doc: Document, page_map: dict[str, int] | None = None):
    page_map = page_map or {}
    toc_idx = next(i for i, p in enumerate(doc.paragraphs) if norm(p.text) == "目录")
    body_start = find_body_start(doc, toc_idx)

    for p in list(doc.paragraphs[toc_idx + 1 : body_start]):
        delete_paragraph(p)

    title = doc.paragraphs[toc_idx]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.style = doc.styles["Report Heading 1"]
    title.paragraph_format.page_break_before = True

    parent = title._p.getparent()
    insert_pos = parent.index(title._p) + 1
    from docx.oxml import OxmlElement

    for level, label, target in TOC_ITEMS:
        p_el = OxmlElement("w:p")
        parent.insert(insert_pos, p_el)
        insert_pos += 1
        p = next(par for par in doc.paragraphs if par._p is p_el)
        p.style = doc.styles["Report TOC"]
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm({1: 0, 2: 0.55, 3: 1.1}[level])
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        page = str(page_map.get(target, ""))
        r1 = p.add_run(label)
        set_run_font(r1, east="宋体", size=10.5 if level == 1 else 10, bold=(level == 1))
        p.add_run("\t")
        r2 = p.add_run(page)
        set_run_font(r2, east="宋体", size=10.5 if level == 1 else 10, bold=(level == 1))

    body_heading = next((p for p in doc.paragraphs if norm(p.text) in {"研究背景与研究问题", "第1章 研究背景与研究问题"}), None)
    if body_heading:
        body_heading.paragraph_format.page_break_before = True


def find_body_start(doc: Document, toc_idx: int) -> int:
    paragraphs = doc.paragraphs
    for i in range(toc_idx + 1, len(paragraphs)):
        if norm(paragraphs[i].text) not in {"研究背景与研究问题", "第1章 研究背景与研究问题"}:
            continue
        following = [norm(p.text) for p in paragraphs[i + 1 : i + 8] if norm(p.text)]
        if "本章xxx" in following:
            return i
    raise RuntimeError("正文起点未找到")


def normalize_chapter_titles(doc: Document):
    replacements = {
        "研究背景与研究问题": "第1章 研究背景与研究问题",
        "研究内容": "第2章 研究内容",
        "原型设计": "第3章 原型设计",
        "实验验证": "第4章 实验验证",
        "下一阶段研究计划": "第5章 下一阶段研究计划",
    }
    seen = set()
    for p in doc.paragraphs:
        t = norm(p.text)
        if t in replacements and t not in seen:
            replace_paragraph_text(p, replacements[t])
            p.style = doc.styles["Report Heading 1"]
            p.paragraph_format.first_line_indent = Cm(0)
            seen.add(t)


def replace_paragraph_text(paragraph, text: str):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run)


def add_model_details(doc: Document):
    targets = {norm(p.text): p for p in doc.paragraphs}

    intro = targets.get("技术方法研究重点包括多智能体语音对话架构、智能体个性化匹配方法和性能优化三部分。")
    if intro:
        replace_paragraph_text(
            intro,
            "技术方法研究重点包括多智能体语音对话架构、智能体个性化匹配方法和性能优化三部分。为保证系统链路可复现，本节同时说明各步骤所使用的模型、调用方式及部署形态。整体上，强意图识别和会话保持采用本地规则，语音画像采用本地部署模型，弱意图语义判断采用大模型 API，最终语音智能体回复则调用 Coze 实时语音智能体。",
        )

    p189 = targets.get("在本研究的系统架构中，用户说出“豆包”“孙悟空”等明确唤醒词时，系统直接调起对应智能体；当用户说“我今天肚子好痛”“我想学英语”这类非唤醒词表达时，系统进入个性化匹配流程。该流程突破了单一唤醒词限制，同时引入会话保持和弱意图路由，减少用户频繁切换智能体带来的打断感。")
    if p189:
        insert_after(
            p189,
            "具体实现上，语音输入在浏览器端优先通过 Web Speech API 转写为文本；当需要后端兜底转写时，调用兼容 OpenAI 接口的 Whisper-1 语音识别模型 API。强意图阶段不调用大模型，而是在本地根据智能体名称、别名和唤醒词进行规则匹配；匹配成功后，系统通过 Coze Realtime SDK/API 调起对应 Coze 语音智能体，使用各智能体在 Coze 中配置的 Bot ID、实时语音访问令牌和音色完成端到端语音对话。",
            "Report Body",
        )

    p198 = targets.get("系统首先通过 ASR 得到用户 query，同时通过语音画像模块识别年龄段和性别；随后将 query 归一到五类场景之一，包括休闲陪伴、健康咨询、智能家居、信息查询和学习辅导。")
    if p198:
        replace_paragraph_text(
            p198,
            "系统首先通过 ASR 得到用户 query，同时通过语音画像模块识别年龄段和性别。语音画像模型采用 Hugging Face 上的 audeering/wav2vec2-large-robust-6-ft-age-gender，即基于 Wav2Vec2 的年龄与性别识别模型，本研究采用本地 snapshot 离线部署，不把用户原始语音上传到第三方画像服务。随后，系统将 query 归一到休闲陪伴、健康咨询、智能家居、信息查询和学习辅导五类场景之一。场景归一优先使用轻量关键词规则处理高置信表达；边界表达则交由 LLM 路由模型判断。",
        )

    p204 = targets.get("该方法的好处在于，数据负责用户偏好召回，LLM 负责语义细节判断。这样既保留了大模型对自然语言的理解能力，也减少了全量候选带来的不稳定和延迟。多模态语音画像模块则用于补充弱意图匹配中的用户特征，只输出年龄段、性别和置信度，不直接决定最终智能体。")
    if p204:
        insert_after(
            p204,
            "LLM 路由模型通过 OpenOneApi 的 OpenAI 兼容接口调用，当前配置模型为 GPT-5.5，采用 Responses API 形式，并将 reasoning effort 控制在较低或最小档以降低等待时间。模型输入包括用户 query、年龄段、性别、当前智能体、Top3 候选智能体及其人设/能力标签；模型只需在 Top3 中输出最终 agentId、匹配理由和置信度。这样可以把大模型用于最需要语义判断的环节，而不是让其承担全部召回任务。",
            "Report Body",
        )

    p209 = targets.get("模型层面，系统由较大的 24 层 wav2vec2 年龄性别识别模型，调整为精简 6 层模型，并固定使用本地 Hugging Face snapshot 离线加载，避免每次联网检索。音频处理层面，系统将输入统一转为 16kHz mono wav，只截取最多约 3 秒有效语音，同时进行前后静音裁剪和音量归一化，减少静音、噪声和音量差异造成的偏差。")
    if p209:
        replace_paragraph_text(
            p209,
            "模型层面，系统最初评估的是 audeering/wav2vec2-large-robust-24-ft-age-gender，即 24 层 Wav2Vec2 年龄性别识别模型；由于本地推理耗时较长，后续调整为 audeering/wav2vec2-large-robust-6-ft-age-gender 精简模型，并固定使用本地 Hugging Face snapshot 离线加载，避免每次联网检索。音频处理层面，系统将输入统一转为 16kHz mono wav，只截取最多约 3 秒有效语音，同时进行前后静音裁剪和音量归一化，减少静音、噪声和音量差异造成的偏差。",
        )


def render_docx(docx_path: Path, out_dir: Path):
    if out_dir.exists():
        shutil.rmtree(out_dir)
    subprocess.run([str(PYTHON), str(RENDER), str(docx_path), "--output_dir", str(out_dir), "--emit_pdf"], check=True, cwd=PROJECT_ROOT)


def extract_page_map(pdf_path: Path) -> dict[str, int]:
    if not PDFTOTEXT.exists():
        return {}
    text_path = TEMP_DIR / "rendered.txt"
    subprocess.run([str(PDFTOTEXT), "-layout", str(pdf_path), str(text_path)], check=True)
    raw = text_path.read_text(errors="ignore")
    pages = raw.split("\f")
    page_map = {}
    targets = [target for _, _, target in TOC_ITEMS]
    for idx, page in enumerate(pages, start=1):
        compact = norm(page)
        if "目录" in compact:
            continue
        for target in targets:
            if target not in page_map and target in compact:
                page_map[target] = idx
    return page_map


def main():
    if not BACKUP.exists():
        shutil.copy2(SRC, BACKUP)

    doc = Document(BACKUP)
    ensure_style(doc, "Report Body", "仿宋_GB2312", 12, False, True)
    ensure_style(doc, "Report Heading 1", "黑体", 15, True)
    ensure_style(doc, "Report Heading 2", "黑体", 13, True)
    ensure_style(doc, "Report Heading 3", "黑体", 12, True)
    ensure_style(doc, "Report Caption", "仿宋_GB2312", 10.5)
    toc_style = ensure_style(doc, "Report TOC", "宋体", 10)
    toc_style.paragraph_format.line_spacing = 1.0
    toc_style.paragraph_format.space_after = Pt(0)

    replace_toc(doc)
    normalize_chapter_titles(doc)
    apply_heading_styles(doc)
    add_model_details(doc)
    doc.save(SRC)

    TEMP_DIR.mkdir(parents=True, exist_ok=True)
    render_docx(SRC, TEMP_DIR / "first_render")
    page_map = extract_page_map(TEMP_DIR / "first_render" / "任科中期报告_0621.pdf")
    if page_map:
        doc = Document(SRC)
        replace_toc(doc, page_map)
        doc.save(SRC)

    print(SRC)
    print(BACKUP)


if __name__ == "__main__":
    main()
